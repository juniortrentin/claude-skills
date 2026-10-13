#!/usr/bin/env python3
"""Confere notas, casa itens com o cadastro e grava a planilha da skill Cupom fiscal vira cadastro.

Trentin Solutions · criada por Ivaldo Trentin · trentinsolutions.com.br
Uso livre, mantendo este crédito (ver LICENSE.md).

Três etapas, nesta ordem:

  python lancar.py conferir trabalho/notas
      Confere cada nota (linha a linha e total) e acha notas repetidas no lote.

  python lancar.py casar trabalho/notas [--base planilha.xlsx | --catalogo produtos.csv --compras compras.csv]
                         [--config trabalho/config.json]
      Casa cada item com o cadastro (EAN, código, descrição), rateia desconto e frete,
      abre kits, calcula custo e preço sugerido. Gera trabalho/proposta.json.

  python lancar.py planilha trabalho/proposta.json --saida destino.xlsx [--base planilha.xlsx]
      Sem --base: cria a planilha modelo já preenchida.
      Com --base no formato modelo: acrescenta as linhas numa cópia.
      Sempre grava também compras.csv, produtos_novos.csv e despesas.csv (para colar no Google Sheets).
"""
import argparse
import csv
import glob
import json
import math
import os
import re
import sys
import unicodedata
from datetime import date
from difflib import SequenceMatcher

CONFIG_PADRAO = {
    "markup": 2.0,
    "arredondar": True,
    "precos": [2, 3, 4, 5, 6, 7, 8, 10, 12, 15, 18, 20, 25, 30, 35, 40, 45, 50, 60, 70, 80, 90, 100,
               120, 150, 180, 200, 250, 300],
    "ratear_desconto": True,
    "ratear_frete": True,
    "tolerancia": 0.05,
}

COL_PRODUTOS = ["Código", "Produto", "EAN", "Fornecedor", "Unidade", "Custo médio", "Preço de venda",
                "Markup", "Qtd comprada", "Última compra", "Status", "Observações"]
COL_COMPRAS = ["ID compra", "Data", "Fornecedor", "Nº nota", "Código", "Produto", "Qtd", "Custo unit.",
               "Custo total", "Chave de acesso", "Origem", "Observações"]
COL_FORN = ["CNPJ", "Fornecedor", "Razão social", "Primeira compra", "Observações"]
COL_DESP = ["Data", "Fornecedor", "Descrição", "Qtd", "Valor", "Nº nota", "Observações"]
COL_CONF = ["ID compra", "Fornecedor", "Nº nota", "Data", "Itens", "Soma dos itens", "Desconto", "Frete e outros",
            "Total da nota", "Diferença", "Status", "Origem"]

CREDITO = [
    "Skill gratuita da Trentin Solutions, criada por Ivaldo Trentin · trentinsolutions.com.br",
    "Esta planilha te ajudou? Um Pix de qualquer valor em agradecimento é bem-vindo.",
    "Chave Pix aleatória: b73301d4-682f-46a0-857a-3a7653c2befb · Ivaldo Idalgo Trentin Junior Ltda",
]


# ---------- utilidades ----------
def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def parecido(a, b):
    return SequenceMatcher(None, norm(a), norm(b)).ratio()


def sigla(nome):
    palavras = [p for p in norm(nome).split() if p not in ("de", "da", "do", "e", "ltda", "me", "eireli", "sa", "comercio")]
    s = "".join(p[0] for p in palavras[:3]).upper() if len(palavras) >= 2 else (palavras[0][:3].upper() if palavras else "FOR")
    return s or "FOR"


def f2(x):
    return round(float(x or 0) + 1e-9, 2)


def para_data(s):
    """'2026-10-05' -> date, para que filtros, MAXIFS e formatação de data funcionem na planilha."""
    try:
        return date.fromisoformat(str(s)[:10])
    except ValueError:
        return s


def preco_sugerido(custo, cfg):
    bruto = custo * float(cfg["markup"])
    if not cfg.get("arredondar"):
        return round(bruto, 2)
    for p in cfg["precos"]:
        if p >= bruto - 0.009:
            return float(p)
    return float(math.ceil(bruto / 10.0) * 10)


def carregar_config(caminho):
    cfg = dict(CONFIG_PADRAO)
    if caminho and os.path.exists(caminho):
        with open(caminho, encoding="utf-8") as f:
            cfg.update(json.load(f))
    return cfg


def carregar_notas(pasta):
    notas = []
    for arq in sorted(glob.glob(os.path.join(pasta, "*.json"))):
        with open(arq, encoding="utf-8") as f:
            try:
                n = json.load(f)
            except json.JSONDecodeError as e:
                print(f"ERRO {arq}: JSON inválido ({e})", file=sys.stderr)
                continue
        n["_arquivo"] = os.path.basename(arq)
        notas.append(n)
    return notas


def nome_forn(n):
    f = n.get("fornecedor") or {}
    return (f.get("nome") or f.get("razao_social") or "Fornecedor sem nome").strip()


def chave_nota(n):
    if re.fullmatch(r"\d{44}", str(n.get("chave") or "")):
        return "CH:" + n["chave"]
    f = n.get("fornecedor") or {}
    quem = re.sub(r"\D", "", f.get("cnpj") or "") or norm(nome_forn(n))
    return f"NF:{quem}|{n.get('numero', '')}|{n.get('data', '')}|{f2(n.get('total'))}"


# ---------- leitura de planilhas existentes ----------
ALIASES = {
    "codigo": ["codigo", "cod", "sku", "referencia", "ref"],
    "produto": ["produto", "descricao", "nome", "item"],
    "ean": ["ean", "gtin", "codigo de barras"],
    "fornecedor": ["fornecedor", "loja", "marca"],
    "preco": ["preco de venda", "preco venda", "preco", "valor de venda"],
    "chave": ["chave de acesso", "chave"],
    "nota": ["n nota", "no nota", "nota", "numero da nota", "n nota cupom", "nota cupom"],
}


def mapear(cabecalho):
    cab = [norm(c) for c in cabecalho]
    m = {}
    for campo, nomes in ALIASES.items():
        for i, c in enumerate(cab):
            if c in nomes or any(c.startswith(n) for n in nomes):
                m.setdefault(campo, i)
    return m


def linhas_de(caminho, aba_preferida):
    if caminho.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(caminho, data_only=True, read_only=True)
        nomes = {norm(s): s for s in wb.sheetnames}
        aba = nomes.get(norm(aba_preferida))
        if not aba:
            return []
        return [list(r) for r in wb[aba].iter_rows(values_only=True)]
    with open(caminho, encoding="utf-8-sig") as f:
        return list(csv.reader(f))


def ler_catalogo(caminho):
    linhas = linhas_de(caminho, "Produtos") if caminho else []
    if not linhas:
        return []
    m = mapear(linhas[0])
    if "produto" not in m:
        print("Aviso: não achei a coluna de nome do produto no cadastro; tudo será tratado como novo.", file=sys.stderr)
        return []
    cat = []
    for r in linhas[1:]:
        g = lambda k: (str(r[m[k]]).strip() if k in m and m[k] < len(r) and r[m[k]] is not None else "")
        if g("produto") or g("codigo"):
            cat.append({"codigo": g("codigo").lstrip("'"), "produto": g("produto"), "ean": g("ean"),
                        "fornecedor": g("fornecedor"), "preco": g("preco")})
    return cat


def ler_notas_lancadas(caminho):
    linhas = linhas_de(caminho, "Compras") if caminho else []
    if not linhas:
        return set(), set()
    m = mapear(linhas[0])
    chaves, notas = set(), set()
    for r in linhas[1:]:
        g = lambda k: (str(r[m[k]]).strip() if k in m and m[k] < len(r) and r[m[k]] is not None else "")
        if re.fullmatch(r"\d{44}", g("chave")):
            chaves.add(g("chave"))
        if g("nota") and g("fornecedor"):
            notas.add((norm(g("fornecedor")), g("nota").lstrip("'")))
    return chaves, notas


# ---------- etapa 1: conferir ----------
def conferir_nota(n, tol):
    problemas = []
    itens = n.get("itens") or []
    if not itens:
        problemas.append("nota sem itens")
    soma = 0.0
    for i, it in enumerate(itens, 1):
        q, u, t = float(it.get("qtd") or 0), float(it.get("valor_unit") or 0), float(it.get("valor_total") or 0)
        if q <= 0:
            problemas.append(f"item {i} sem quantidade")
        if t == 0 and q and u:
            it["valor_total"] = t = round(q * u, 2)
        if q and u and abs(q * u - t) > max(0.02, t * 0.005):
            problemas.append(f"item {i} ({it.get('descricao', '')[:30]}): {q} × {u:.2f} = {q * u:.2f}, impresso {t:.2f}")
        soma += t - float(it.get("desconto") or 0) + float(it.get("frete") or 0) + float(it.get("outros") or 0)
    desc, frete, outros = float(n.get("desconto") or 0), float(n.get("frete") or 0), float(n.get("outros") or 0)
    calculado = round(soma - desc + frete + outros, 2)
    total = float(n.get("total") or 0)
    dif = round(total - calculado, 2) if total else 0.0
    if not total:
        problemas.append("total da nota não informado")
    elif abs(dif) > tol:
        problemas.append(f"soma dos itens {calculado:.2f} ≠ total {total:.2f} (diferença {dif:+.2f})")
    ch = str(n.get("chave") or "")
    if ch:
        sys.path.insert(0, os.path.dirname(__file__))
        from ler_qrcode import chave_valida
        if not chave_valida(ch):
            problemas.append("chave de acesso com dígito verificador inválido (confira a leitura)")
    return {"soma_itens": round(soma, 2), "calculado": calculado, "diferenca": dif, "problemas": problemas}


def cmd_conferir(a):
    cfg = carregar_config(a.config)
    notas = carregar_notas(a.pasta)
    vistos, saida = {}, []
    for n in notas:
        r = conferir_nota(n, cfg["tolerancia"])
        k = chave_nota(n)
        status = "ok" if not r["problemas"] else "conferir"
        if k in vistos:
            status = "repetida"
            r["problemas"].append(f"mesma nota que {vistos[k]}")
        else:
            vistos[k] = n["_arquivo"]
        saida.append({"arquivo": n["_arquivo"], "fornecedor": nome_forn(n), "numero": n.get("numero", ""),
                      "data": n.get("data", ""), "itens": len(n.get("itens") or []), "total": f2(n.get("total")),
                      "status": status, **r})
    with open(os.path.join(os.path.dirname(a.pasta.rstrip("/")) or ".", "conferencia.json"), "w", encoding="utf-8") as f:
        json.dump(saida, f, ensure_ascii=False, indent=2)
    forn = len({s["fornecedor"] for s in saida})
    print(f"{len(saida)} notas · {forn} fornecedores")
    for s in saida:
        marca = {"ok": "OK      ", "conferir": "CONFERIR", "repetida": "REPETIDA"}[s["status"]]
        print(f"{marca} {s['fornecedor'][:28]:<28} nº {str(s['numero'])[:10]:<10} {s['data']}  {s['itens']:>3} itens  R$ {s['total']:>10.2f}")
        for p in s["problemas"]:
            print(f"         - {p}")


# ---------- etapa 2: casar ----------
def casar_item(it, forn, cat, novos_lote):
    ean, cod, desc = str(it.get("ean") or ""), str(it.get("codigo") or "").strip(), it.get("descricao", "")
    if ean:
        for p in cat:
            if p["ean"] and p["ean"] == ean:
                return "reposicao", p, "mesmo EAN"
    if cod:
        for p in cat:
            if p["codigo"] == cod and (norm(p["fornecedor"]) == norm(forn) or parecido(p["produto"], desc) >= 0.8):
                return "reposicao", p, "mesmo código"
    melhor, nota_ = None, 0.0
    for p in cat:
        r = parecido(p["produto"], desc)
        if norm(p["fornecedor"]) != norm(forn):
            r -= 0.05
        if r > nota_:
            melhor, nota_ = p, r
    if melhor and nota_ >= 0.92:
        return "reposicao", melhor, f"descrição parecida ({nota_:.0%})"
    chave_lote = (norm(forn), cod or norm(desc))
    if chave_lote in novos_lote:
        return "novo_repetido", novos_lote[chave_lote], "já apareceu neste lote"
    if melhor and nota_ >= 0.8:
        return "novo", melhor, f"parecido com {melhor['codigo']} ({nota_:.0%}): conferir se não é o mesmo"
    return "novo", None, ""


def cmd_casar(a):
    cfg = carregar_config(a.config)
    cat = ler_catalogo(a.base or a.catalogo)
    chaves_ja, notas_ja = ler_notas_lancadas(a.base or a.compras)
    conf = {}
    conf_arq = os.path.join(os.path.dirname(a.pasta.rstrip("/")) or ".", "conferencia.json")
    if os.path.exists(conf_arq):
        with open(conf_arq, encoding="utf-8") as f:
            conf = {c["arquivo"]: c for c in json.load(f)}
    codigos_usados = {p["codigo"] for p in cat if p["codigo"]}
    prov_seq = {}
    novos_lote, compras, novos, despesas, forns, resumo, puladas = {}, [], [], [], {}, [], []
    vistos = set()
    for n in carregar_notas(a.pasta):
        forn = nome_forn(n)
        c = conf.get(n["_arquivo"]) or {"status": "ok", "problemas": []}
        k = chave_nota(n)
        motivo = ""
        if k in vistos or c["status"] == "repetida":
            motivo = "repetida no lote"
        elif n.get("chave") in chaves_ja or (norm(forn), str(n.get("numero", ""))) in notas_ja:
            motivo = "já lançada na planilha"
        elif c["status"] == "conferir" and not a.incluir_divergentes and not n.get("liberada"):
            motivo = "divergência: " + "; ".join(c["problemas"])
        if motivo:
            puladas.append({"arquivo": n["_arquivo"], "fornecedor": forn, "numero": n.get("numero", ""), "motivo": motivo})
            continue
        vistos.add(k)
        sg = sigla(forn)
        id_compra = f"{sg}-{str(n.get('data', ''))[:10].replace('-', '')}-{n.get('numero') or len(resumo) + 1}"
        cnpj = re.sub(r"\D", "", (n.get("fornecedor") or {}).get("cnpj") or "")
        forns.setdefault(cnpj or norm(forn), {"cnpj": cnpj, "nome": forn,
                                             "razao_social": (n.get("fornecedor") or {}).get("razao_social", ""),
                                             "primeira": n.get("data", "")})
        itens = n.get("itens") or []
        base = [float(i.get("valor_total") or 0) - float(i.get("desconto") or 0) for i in itens]
        extra = 0.0
        if cfg["ratear_desconto"]:
            extra -= float(n.get("desconto") or 0)
        if cfg["ratear_frete"]:
            extra += float(n.get("frete") or 0) + float(n.get("outros") or 0)
        soma_base = sum(base) or 1.0
        origem = n.get("origem", "")
        for it, b in zip(itens, base):
            custo_total = b + float(it.get("frete") or 0) + float(it.get("outros") or 0) + extra * b / soma_base
            fator = float(it.get("fator") or 1) or 1
            qtd = float(it.get("qtd") or 0) * fator
            custo_unit = custo_total / qtd if qtd else 0.0
            obs = []
            if fator != 1:
                obs.append(f"kit aberto: {it.get('qtd')} × {fator:g} un")
            if it.get("descricao_original"):
                obs.append(f'na nota: "{it["descricao_original"]}"')
            if it.get("tipo") == "despesa":
                despesas.append({"data": n.get("data", ""), "fornecedor": forn, "descricao": it.get("descricao", ""),
                                 "qtd": qtd, "valor": f2(custo_total), "nota": n.get("numero", ""),
                                 "obs": "; ".join(obs)})
                continue
            status, ref, porque = casar_item(it, forn, cat, novos_lote)
            if status == "reposicao":
                codigo, nome = ref["codigo"], ref["produto"]
                obs.append(f"reposição ({porque})")
            elif status == "novo_repetido":
                codigo, nome = ref["codigo"], ref["produto"]
            else:
                cod = str(it.get("codigo") or "").strip()
                if cod and cod not in codigos_usados:
                    codigo = cod
                elif cod:
                    codigo = f"{sg}-{cod}"
                else:
                    prov_seq[sg] = prov_seq.get(sg, 0) + 1
                    while f"{sg}-{prov_seq[sg]:03d}" in codigos_usados:
                        prov_seq[sg] += 1
                    codigo = f"{sg}-{prov_seq[sg]:03d}"
                codigos_usados.add(codigo)
                nome = it.get("descricao", "").strip()
                preco = preco_sugerido(custo_unit, cfg)
                obs_p = [f"preço sugerido (markup {float(cfg['markup']):g}×)"]
                if not str(it.get("codigo") or "").strip():
                    obs_p.append("código provisório")
                if porque:
                    obs_p.append(porque)
                p = {"codigo": codigo, "produto": nome, "ean": it.get("ean", ""), "fornecedor": forn,
                     "unidade": "UN" if fator != 1 else it.get("unidade", ""), "preco": preco, "obs": "; ".join(obs_p)}
                novos.append(p)
                novos_lote[(norm(forn), str(it.get("codigo") or "").strip() or norm(nome))] = {"codigo": codigo, "produto": nome}
                obs.append("produto novo")
            compras.append({"id": id_compra, "data": n.get("data", ""), "fornecedor": forn, "nota": n.get("numero", ""),
                            "codigo": codigo, "produto": nome, "qtd": qtd, "custo_unit": round(custo_unit, 4),
                            "custo_total": f2(custo_total), "chave": n.get("chave", ""), "origem": origem,
                            "obs": "; ".join(obs)})
        resumo.append({"id": id_compra, "fornecedor": forn, "numero": n.get("numero", ""), "data": n.get("data", ""),
                       "itens": len(itens), "soma_itens": c.get("soma_itens", f2(sum(base))),
                       "desconto": f2(n.get("desconto")), "frete": f2(float(n.get("frete") or 0) + float(n.get("outros") or 0)),
                       "total": f2(n.get("total")), "diferenca": c.get("diferenca", 0.0),
                       "status": "lançada" if c["status"] == "ok" else "lançada com divergência", "origem": origem})
    proposta = {"gerado_em": date.today().isoformat(), "config": cfg, "notas": resumo, "compras": compras,
                "produtos_novos": novos, "despesas": despesas, "fornecedores": list(forns.values()), "puladas": puladas}
    destino = a.saida or os.path.join(os.path.dirname(a.pasta.rstrip("/")) or ".", "proposta.json")
    with open(destino, "w", encoding="utf-8") as f:
        json.dump(proposta, f, ensure_ascii=False, indent=2)
    repos = sum(1 for c in compras if "reposição" in c["obs"])
    print(f"{len(resumo)} notas prontas · {len(compras)} linhas de compra ({repos} reposições, {len(novos)} produtos novos)"
          f" · {len(despesas)} despesas · {len(puladas)} notas puladas -> {destino}")
    for p in puladas:
        print(f"  pulada: {p['fornecedor']} nº {p['numero']} — {p['motivo']}")
    duvidas = [p for p in novos if "conferir" in p["obs"]]
    for p in duvidas:
        print(f"  conferir: {p['codigo']} {p['produto'][:40]} — {p['obs']}")


# ---------- etapa 3: planilha ----------
def estilo_cabecalho(ws, n):
    from openpyxl.styles import Alignment, Font, PatternFill
    for c in ws[1][:n]:
        c.fill = PatternFill("solid", fgColor="3A6AB4")
        c.font = Font(bold=True, color="FFFFFF", name="Arial")
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def nova_planilha(cfg):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.properties.creator = "Trentin Solutions · Cupom fiscal vira cadastro"
    wb.properties.description = "Gerada pela skill gratuita Cupom fiscal vira cadastro, da Trentin Solutions (trentinsolutions.com.br)."
    abas = [("Produtos", COL_PRODUTOS, [12, 44, 16, 24, 9, 13, 14, 10, 13, 14, 14, 48]),
            ("Compras", COL_COMPRAS, [20, 12, 24, 12, 12, 44, 8, 12, 12, 48, 10, 48]),
            ("Fornecedores", COL_FORN, [18, 28, 36, 16, 40]),
            ("Despesas", COL_DESP, [12, 24, 44, 8, 12, 12, 40]),
            ("Conferência", COL_CONF, [20, 28, 12, 12, 8, 14, 12, 14, 14, 12, 22, 10])]
    for i, (nome, cols, larg) in enumerate(abas):
        ws = wb.active if i == 0 else wb.create_sheet(nome)
        ws.title = nome
        ws.append(cols)
        for j, w in enumerate(larg, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        estilo_cabecalho(ws, len(cols))
    wc = wb.create_sheet("Config")
    wc.append(["Configuração", "Valor"])
    estilo_cabecalho(wc, 2)
    for k in ("markup", "arredondar", "ratear_desconto", "ratear_frete", "tolerancia"):
        wc.append([k, cfg[k]])
    wc.append(["precos", ", ".join(f"{p:g}" for p in cfg["precos"])])
    wc.column_dimensions["A"].width = 20
    wc.column_dimensions["B"].width = 90
    wu = wb.create_sheet("Como usar")
    texto = [
        "Cupom fiscal vira cadastro · trentinsolutions.com.br", "",
        "Produtos: um produto por linha. Custo médio, Markup, Qtd comprada e Última compra são fórmulas: não digite nelas.",
        "Compras: uma linha por item de nota. Custo total é fórmula (Qtd × Custo unit.).",
        "O custo médio é ponderado: soma de todos os custos do produto ÷ soma das quantidades compradas.",
        "Status 'Novo — conferir' some quando você revisar o produto e apagar o texto.",
        "Preço de venda sugerido = custo × markup (aba Config), arredondado para um preço comum.",
        "Conferência: cada nota com a soma dos itens comparada ao total impresso.", "",
        *CREDITO,
    ]
    for t in texto:
        wu.append([t])
    wu["A1"].font = Font(name="Arial", size=13, bold=True, color="3A6AB4")
    wu.column_dimensions["A"].width = 120
    return wb


def e_modelo(wb):
    return "Produtos" in wb.sheetnames and "Compras" in wb.sheetnames and \
        [c.value for c in wb["Compras"][1][:len(COL_COMPRAS)]] == COL_COMPRAS


def ultima_linha(ws):
    r = ws.max_row
    while r > 1 and all(c.value in (None, "") for c in ws[r][:3]):
        r -= 1
    return r


def gravar_xlsx(prop, saida, base):
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    if base:
        wb = load_workbook(base)
        if not e_modelo(wb):
            return False
    else:
        wb = nova_planilha(prop["config"])
    wp, wc, wf, wd, wk = (wb[n] for n in ("Produtos", "Compras", "Fornecedores", "Despesas", "Conferência"))
    corpo = Font(name="Arial", size=10)

    r = ultima_linha(wc)
    for c in prop["compras"]:
        r += 1
        vals = [c["id"], para_data(c["data"]), c["fornecedor"], str(c["nota"]), str(c["codigo"]), c["produto"], c["qtd"],
                c["custo_unit"], f"=G{r}*H{r}", c["chave"], c["origem"], c["obs"]]
        for j, v in enumerate(vals, 1):
            cel = wc.cell(row=r, column=j, value=v)
            cel.font = corpo
        wc.cell(row=r, column=2).number_format = "DD/MM/YYYY"
        wc.cell(row=r, column=8).number_format = '"R$" #,##0.00'
        wc.cell(row=r, column=9).number_format = '"R$" #,##0.00'

    r = ultima_linha(wp)
    for p in prop["produtos_novos"]:
        r += 1
        vals = [str(p["codigo"]), p["produto"], p["ean"], p["fornecedor"], p["unidade"],
                f'=IFERROR(SUMIFS(Compras!I:I,Compras!E:E,A{r})/SUMIFS(Compras!G:G,Compras!E:E,A{r}),0)',
                p["preco"], f'=IFERROR(G{r}/F{r},"")', f"=SUMIFS(Compras!G:G,Compras!E:E,A{r})",
                f'=IFERROR(_xlfn.MAXIFS(Compras!B:B,Compras!E:E,A{r}),"")', "Novo — conferir", p["obs"]]
        for j, v in enumerate(vals, 1):
            wp.cell(row=r, column=j, value=v).font = corpo
        for j in (6, 7):
            wp.cell(row=r, column=j).number_format = '"R$" #,##0.00'
        wp.cell(row=r, column=8).number_format = '0.00"×"'
        wp.cell(row=r, column=10).number_format = "DD/MM/YYYY"

    existentes = {str(row[0].value or "") or norm(row[1].value or "") for row in wf.iter_rows(min_row=2) if row[0].value or row[1].value}
    r = ultima_linha(wf)
    for f_ in prop["fornecedores"]:
        k = f_["cnpj"] or norm(f_["nome"])
        if k in existentes:
            continue
        r += 1
        for j, v in enumerate([f_["cnpj"], f_["nome"], f_["razao_social"], para_data(f_["primeira"]), ""], 1):
            wf.cell(row=r, column=j, value=v).font = corpo

    r = ultima_linha(wd)
    for d in prop["despesas"]:
        r += 1
        for j, v in enumerate([para_data(d["data"]), d["fornecedor"], d["descricao"], d["qtd"], d["valor"], str(d["nota"]), d["obs"]], 1):
            wd.cell(row=r, column=j, value=v).font = corpo

    r = ultima_linha(wk)
    for n in prop["notas"]:
        r += 1
        vals = [n["id"], n["fornecedor"], str(n["numero"]), para_data(n["data"]), n["itens"], n["soma_itens"], n["desconto"],
                n["frete"], n["total"], n["diferenca"], n["status"], n["origem"]]
        for j, v in enumerate(vals, 1):
            wk.cell(row=r, column=j, value=v).font = corpo

    wb.save(saida)
    return True


def gravar_csvs(prop, pasta):
    os.makedirs(pasta, exist_ok=True)
    def w(nome, cab, linhas):
        with open(os.path.join(pasta, nome), "w", newline="", encoding="utf-8-sig") as f:
            cw = csv.writer(f)
            cw.writerow(cab)
            cw.writerows(linhas)
    w("compras.csv", [c for c in COL_COMPRAS if c != "Custo total"] + ["Custo total (calculado)"],
      [[c["id"], c["data"], c["fornecedor"], c["nota"], c["codigo"], c["produto"], c["qtd"], c["custo_unit"],
        c["chave"], c["origem"], c["obs"], c["custo_total"]] for c in prop["compras"]])
    w("produtos_novos.csv", ["Código", "Produto", "EAN", "Fornecedor", "Unidade", "Preço de venda sugerido", "Observações"],
      [[p["codigo"], p["produto"], p["ean"], p["fornecedor"], p["unidade"], p["preco"], p["obs"]] for p in prop["produtos_novos"]])
    w("despesas.csv", COL_DESP,
      [[d["data"], d["fornecedor"], d["descricao"], d["qtd"], d["valor"], d["nota"], d["obs"]] for d in prop["despesas"]])


def cmd_planilha(a):
    with open(a.proposta, encoding="utf-8") as f:
        prop = json.load(f)
    pasta_csv = os.path.join(os.path.dirname(os.path.abspath(a.saida)), "csv")
    gravar_csvs(prop, pasta_csv)
    try:
        ok = gravar_xlsx(prop, a.saida, a.base)
    except ImportError:
        print("openpyxl não instalado (pip install openpyxl); gerei só os CSV em", pasta_csv)
        return
    if not ok:
        print(f"A planilha {a.base} não está no formato modelo. Nada foi alterado nela.")
        print(f"Use os CSV em {pasta_csv} e o mapa de colunas para acrescentar as linhas na estrutura dela.")
        return
    print(f"Planilha gravada: {a.saida} · CSV para colar no Google Sheets: {pasta_csv}")
    print(f"{len(prop['notas'])} notas · {len(prop['compras'])} linhas de compra · {len(prop['produtos_novos'])} produtos novos · {len(prop['despesas'])} despesas")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s1 = sub.add_parser("conferir")
    s1.add_argument("pasta")
    s1.add_argument("--config")
    s2 = sub.add_parser("casar")
    s2.add_argument("pasta")
    s2.add_argument("--base", help="planilha .xlsx existente (lê Produtos e Compras)")
    s2.add_argument("--catalogo", help="CSV da aba de produtos (ex.: exportado do Google Sheets)")
    s2.add_argument("--compras", help="CSV da aba de compras, para não lançar nota repetida")
    s2.add_argument("--config")
    s2.add_argument("--saida")
    s2.add_argument("--incluir-divergentes", action="store_true", help="lança também notas que não fecharam")
    s3 = sub.add_parser("planilha")
    s3.add_argument("proposta")
    s3.add_argument("--saida", required=True)
    s3.add_argument("--base")
    a = ap.parse_args()
    {"conferir": cmd_conferir, "casar": cmd_casar, "planilha": cmd_planilha}[a.cmd](a)


if __name__ == "__main__":
    main()
