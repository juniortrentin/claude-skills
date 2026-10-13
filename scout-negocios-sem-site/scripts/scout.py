#!/usr/bin/env python3
"""Scout de negócios sem site: nota de presença digital, oportunidade, ranking de nichos,
planilha e diagnóstico de 1 página por negócio.

Trentin Solutions · criada por Ivaldo Trentin · trentinsolutions.com.br
Uso livre, mantendo este crédito (ver LICENSE.md).

Uso:
  python scout.py nichos trabalho/negocios.jsonl
      Ranking dos nichos: quantos negócios, quantos sem site, oportunidade média.
  python scout.py gerar trabalho/negocios.jsonl --briefing trabalho/briefing.json --marca trabalho/marca.json \
         --saida scout/ [--limite 30] [--diagnosticos 10]
      Planilha (Oportunidades, Nichos, Como usar) e diagnósticos HTML dos mais quentes.

Formato de cada negócio em references/campos.md.
"""
import argparse
import html
import json
import math
import os
import re
import sys
import unicodedata
from datetime import date, datetime
from urllib.parse import quote

HOJE = date.today()

PESOS_OFERTA = {
    "site": {"site": 1.0}, "landing": {"site": 1.0}, "loja": {"site": 0.7, "social": 0.3},
    "google": {"google": 1.0}, "redes": {"social": 1.0}, "trafego": {"site": 0.5, "google": 0.5},
    "whatsapp": {"whats": 1.0},
}
NOME_OFERTA = {"site": "site", "landing": "landing page", "loja": "loja virtual", "google": "perfil no Google",
               "redes": "gestão de redes sociais", "trafego": "tráfego pago", "whatsapp": "atendimento no WhatsApp"}


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", s)).strip()


def slug(s):
    return norm(s).replace(" ", "-")[:60] or "negocio"


def dias_desde(d):
    try:
        return (HOJE - datetime.strptime(str(d)[:10], "%Y-%m-%d").date()).days
    except ValueError:
        return None


def ler(caminho):
    out, vistos = [], set()
    with open(caminho, encoding="utf-8") as f:
        for n, l in enumerate(f, 1):
            l = l.strip()
            if not l:
                continue
            try:
                d = json.loads(l)
            except json.JSONDecodeError as e:
                print(f"Aviso: linha {n} ignorada ({e})", file=sys.stderr)
                continue
            k = (norm(d.get("nome")), re.sub(r"\D", "", str(d.get("telefone") or "")) or norm(d.get("endereco")))
            if k in vistos:
                continue
            vistos.add(k)
            out.append(d)
    return out


# ------------------------------------------------------------ notas
def presenca(n):
    """Nota de presença digital 0-100 e lista de itens do diagnóstico."""
    itens = []
    st = n.get("site_status") or ("ok" if n.get("site") else "sem_site")
    # Site: até 45
    if st == "ok":
        s = 25
        ok_https, ok_mob = bool(n.get("https")), bool(n.get("mobile"))
        ano = n.get("ano_rodape")
        atual = isinstance(ano, int) and ano >= HOJE.year - 1
        wz = bool(n.get("whatsapp_no_site"))
        s += 5 * ok_https + 5 * ok_mob + 5 * atual + 5 * wz
        itens.append(("site", "Site próprio", True, "Tem site e ele abre normalmente."))
        itens.append(("site", "Site seguro (https)", ok_https, "O site abre com cadeado de segurança." if ok_https else "Sem https, o navegador mostra 'não seguro' para quem entra."))
        itens.append(("site", "Abre bem no celular", ok_mob, "O site se ajusta à tela do celular." if ok_mob else "O site não se ajusta ao celular, onde acontece a maioria das buscas por negócios locais."))
        itens.append(("site", "Site atualizado", atual, f"O rodapé indica {ano}." if ano else "Sem sinal de atualização recente."))
        itens.append(("site", "Botão de WhatsApp no site", wz, "Quem chega pelo site encontra o WhatsApp." if wz else "Quem chega pelo site não encontra um botão para chamar no WhatsApp."))
    else:
        s = {"so_rede_social": 5, "erro": 5, "construtor_basico": 15}.get(st, 0)
        txt = {"sem_site": "Quem procura no Google não encontra uma página própria do negócio.",
               "so_rede_social": "O link de 'site' leva só para uma rede social ou página de links.",
               "erro": "O endereço do site não abre ou está com erro.",
               "construtor_basico": "Existe uma página simples de construtor gratuito, com pouco conteúdo."}.get(st, "")
        itens.append(("site", "Site próprio", False, txt))
    site_pts = s
    # Google: até 30
    g = 0
    tem_g = bool(n.get("maps_url") or n.get("avaliacoes"))
    g += 10 * tem_g
    hor = bool(n.get("horario"))
    fotos = int(n.get("fotos_maps") or 0) >= 10
    resp = bool(n.get("responde_avaliacoes"))
    aval = int(n.get("avaliacoes") or 0) >= 20
    g += 5 * hor + 5 * fotos + 5 * resp + 5 * aval
    itens.append(("google", "Perfil no Google Maps", tem_g, "O negócio aparece no Maps, onde o cliente do bairro procura primeiro." if tem_g else "O negócio não aparece no Maps, onde o cliente do bairro procura primeiro."))
    if tem_g:
        itens.append(("google", "Horário de funcionamento", hor, "O horário está cadastrado no Google." if hor else "Sem horário no Google, o cliente não sabe se pode ir agora."))
        itens.append(("google", "Fotos no perfil (10 ou mais)", fotos, f"{n.get('fotos_maps') or 0} fotos hoje."))
        itens.append(("google", "Responde às avaliações", resp, "As avaliações recebem resposta." if resp else "As avaliações ficam sem resposta. Responder mostra cuidado e ajuda na posição no Maps."))
    # Social: até 15
    d = dias_desde(n.get("insta_ultimo_post"))
    if n.get("instagram") and d is not None and d <= 30:
        so, ok_so, txt = 15, True, f"Último post há {d} dias."
    elif n.get("instagram"):
        so, ok_so, txt = 7, False, f"Último post há {d} dias." if d is not None else "Perfil existe, mas sem posts recentes."
    else:
        so, ok_so, txt = 0, False, "Nenhum perfil encontrado."
    itens.append(("social", "Instagram ativo", ok_so, txt))
    # WhatsApp: até 10
    wpp = bool(n.get("whatsapp"))
    itens.append(("whats", "WhatsApp divulgado", wpp, "O WhatsApp aparece no Google ou nas redes." if wpp else "Não encontramos um WhatsApp divulgado no Google nem nas redes."))
    total = min(100, site_pts + g + so + 10 * wpp)
    return total, {"site": site_pts / 45, "google": g / 30, "social": so / 15, "whats": 1.0 * wpp}, itens, st


def forca(n):
    av = int(n.get("avaliacoes") or 0)
    nota = float(n.get("nota") or 0)
    vol = min(1.0, math.log10(av + 1) / math.log10(300))
    q = 0.3 if nota >= 4.5 else 0.2 if nota >= 4.0 else 0.1 if nota else 0.05
    return round(vol * 0.7 + q, 3)


def oportunidade(n, ofertas):
    pres, partes, itens, st = presenca(n)
    gaps = []
    for o in ofertas or ["site"]:
        pesos = PESOS_OFERTA.get(o, {"site": 1.0})
        gaps.append(sum(w * (1 - partes[k]) for k, w in pesos.items()) / sum(pesos.values()))
    gap = max(gaps) if gaps else 0
    f = forca(n)
    # negócio fraco com tudo faltando não deve passar na frente de negócio forte com lacuna clara
    op = round(100 * (gap ** 0.8) * (0.2 + 0.8 * f))
    return {"presenca": pres, "oportunidade": op, "forca": f, "gap": round(gap, 3), "itens": itens, "site_status": st,
            "prioridade": "Quente" if op >= 70 else "Morno" if op >= 45 else "Frio"}


def melhorias(r, ofertas):
    falhas = [i for i in r["itens"] if not i[2]]
    ordem = {"site": 0, "google": 1, "whats": 2, "social": 3}
    foco = {k for o in ofertas for k in PESOS_OFERTA.get(o, {})}
    falhas.sort(key=lambda i: (0 if i[0] in foco else 1, ordem.get(i[0], 9)))
    return falhas[:3]


# ------------------------------------------------------------ saídas
def tabela_nichos(regs):
    grupos = {}
    for r in regs:
        g = grupos.setdefault(r.get("nicho") or "Outros", [])
        g.append(r)
    linhas = []
    for nicho, rs in grupos.items():
        sem = sum(1 for r in rs if r["_r"]["site_status"] in ("sem_site", "so_rede_social", "erro"))
        linhas.append({"nicho": nicho, "total": len(rs), "sem_site": sem, "pct": sem / len(rs),
                       "quentes": sum(1 for r in rs if r["_r"]["prioridade"] == "Quente"),
                       "media_aval": round(sum(int(r.get("avaliacoes") or 0) for r in rs) / len(rs)),
                       "op_media": round(sum(r["_r"]["oportunidade"] for r in rs) / len(rs))})
    linhas.sort(key=lambda x: (-x["op_media"], -x["sem_site"]))
    return linhas


def cmd_nichos(a):
    regs = ler(a.entrada)
    for r in regs:
        r["_r"] = oportunidade(r, a.ofertas.split(","))
    print(f"{len(regs)} negócios levantados")
    print(f"{'Nicho':<28}{'Total':>6}{'Sem site':>10}{'%':>6}{'Quentes':>9}{'Aval. média':>13}{'Oport.':>8}")
    for l in tabela_nichos(regs):
        print(f"{l['nicho'][:27]:<28}{l['total']:>6}{l['sem_site']:>10}{l['pct']:>6.0%}{l['quentes']:>9}{l['media_aval']:>13}{l['op_media']:>8}")


def gravar_xlsx(regs, nichos, saida, briefing):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    azul, branco = PatternFill("solid", fgColor="3A6AB4"), Font(bold=True, color="FFFFFF", name="Arial")
    corpo = Font(name="Arial", size=10)
    cores = {"Quente": "FFE1D6", "Morno": "FCEFD6", "Frio": "ECEFF4"}
    wb = Workbook()
    wb.properties.creator = "Trentin Solutions · Scout de negócios sem site"
    wb.properties.description = "Gerada pela skill gratuita Scout de negócios sem site, da Trentin Solutions (trentinsolutions.com.br)."
    ws = wb.active
    ws.title = "Oportunidades"
    cols = ["Oportunidade", "Prioridade", "Negócio", "Nicho", "Bairro", "Cidade", "Presença digital", "Site", "Situação do site",
            "Nota Google", "Avaliações", "Instagram", "WhatsApp", "Telefone", "O que falta (principal)", "Diagnóstico",
            "Mensagem sugerida", "Google Maps", "Observações"]
    ws.append(cols)
    sit = {"sem_site": "sem site", "so_rede_social": "só rede social", "erro": "com erro", "construtor_basico": "página básica", "ok": "ok"}
    for r in regs:
        x = r["_r"]
        falta = "; ".join(i[1] for i in x["melhorias"])
        ws.append([x["oportunidade"], x["prioridade"], r.get("nome", ""), r.get("nicho", ""), r.get("bairro", ""), r.get("cidade", ""),
                   x["presenca"], r.get("site", ""), sit.get(x["site_status"], x["site_status"]), r.get("nota", ""),
                   r.get("avaliacoes", ""), r.get("instagram", ""), r.get("whatsapp", ""), r.get("telefone", ""), falta,
                   x.get("diag_arquivo", ""), r.get("mensagem", ""), r.get("maps_url", ""), r.get("observacoes", "")])
    for i, w_ in enumerate([12, 10, 30, 18, 16, 14, 11, 30, 14, 10, 10, 18, 16, 16, 46, 30, 50, 36, 30], 1):
        ws.column_dimensions[get_column_letter(i)].width = w_
    for c in ws[1]:
        c.fill, c.font = azul, branco
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = corpo
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (15, 17, 19))
        row[1].fill = PatternFill("solid", fgColor=cores.get(row[1].value, "FFFFFF"))
        row[0].font = Font(name="Arial", size=10, bold=True)
        for j in (7, 17):
            v = row[j].value
            if v and str(v).startswith("http"):
                row[j].hyperlink = v
                row[j].font = Font(name="Arial", size=10, color="3A6AB4", underline="single")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    wn = wb.create_sheet("Nichos")
    wn.append(["Nicho", "Negócios", "Sem site (ou só rede social)", "% sem site", "Quentes", "Avaliações em média", "Oportunidade média"])
    for l in nichos:
        wn.append([l["nicho"], l["total"], l["sem_site"], l["pct"], l["quentes"], l["media_aval"], l["op_media"]])
    for row in wn.iter_rows(min_row=2):
        row[3].number_format = "0%"
        for c in row:
            c.font = corpo
    for c in wn[1]:
        c.fill, c.font = azul, branco
    for i, w_ in enumerate([28, 10, 24, 11, 9, 18, 18], 1):
        wn.column_dimensions[get_column_letter(i)].width = w_

    wc = wb.create_sheet("Como usar")
    ofertas = ", ".join(NOME_OFERTA.get(o, o) for o in briefing.get("ofertas", ["site"]))
    for t in ["Scout de negócios sem site · trentinsolutions.com.br", "",
              f"Briefing: {briefing.get('cidade', '')} · nichos: {', '.join(briefing.get('nichos', [])) or 'todos'} · oferta: {ofertas}",
              "Presença digital (0-100): site até 45, Google até 30, Instagram até 15, WhatsApp até 10. Quanto menor, mais espaço para melhorar.",
              "Oportunidade (0-100): o que falta naquilo que você vende, multiplicado pela força do negócio (avaliações e nota).",
              "Quente ≥ 70: negócio estabelecido com lacuna clara. Comece por eles.",
              "Diagnóstico: página de 1 folha por negócio, para enviar como abertura da conversa. Confira antes de mandar.",
              "LGPD: dados públicos de empresas, usados para contato comercial. Remova quem pedir para sair.", "",
              "Skill gratuita da Trentin Solutions, criada por Ivaldo Trentin · trentinsolutions.com.br",
              "Esta planilha te ajudou? Um Pix de qualquer valor em agradecimento é bem-vindo.",
              "Chave Pix aleatória: b73301d4-682f-46a0-857a-3a7653c2befb · Ivaldo Idalgo Trentin Junior Ltda"]:
        wc.append([t])
    wc["A1"].font = Font(name="Arial", size=13, bold=True, color="3A6AB4")
    wc.column_dimensions["A"].width = 120
    wb.save(saida)


DIAG = r"""<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Presença digital · __NEGOCIO__</title>
<style>
:root{--cor:__COR__;--tinta:#16181d;--suave:#636874;--linha:#e6e7ea;--ok:#1f8a5b;--falta:#c2412d}
*{box-sizing:border-box}body{margin:0;background:#f4f4f2;color:var(--tinta);font:15px/1.5 system-ui,-apple-system,"Segoe UI",Roboto,sans-serif}
.folha{max-width:780px;margin:24px auto;background:#fff;border-radius:18px;overflow:hidden;box-shadow:0 10px 40px -20px rgba(0,0,0,.25)}
header{background:var(--cor);color:#fff;padding:26px 30px}
header small{opacity:.85;letter-spacing:.06em;text-transform:uppercase;font-size:.72rem}
h1{margin:6px 0 2px;font-size:1.7rem;line-height:1.15}
header p{margin:0;opacity:.9}
.topo{display:grid;grid-template-columns:150px 1fr;gap:24px;padding:26px 30px;border-bottom:1px solid var(--linha);align-items:center}
.gauge{width:150px;height:150px;border-radius:50%;display:grid;place-items:center;background:conic-gradient(var(--cor) calc(var(--v)*1%),#eceef1 0)}
.gauge div{width:116px;height:116px;border-radius:50%;background:#fff;display:grid;place-items:center;text-align:center}
.gauge b{font-size:2.3rem;line-height:1}.gauge span{font-size:.72rem;color:var(--suave)}
.topo h2{margin:0 0 6px;font-size:1.15rem}.topo p{margin:0;color:var(--suave)}
section{padding:22px 30px;border-bottom:1px solid var(--linha)}
h3{margin:0 0 12px;font-size:.78rem;letter-spacing:.08em;text-transform:uppercase;color:var(--suave)}
ul{list-style:none;margin:0;padding:0;display:grid;gap:10px}
ul li{display:grid;grid-template-columns:26px 1fr;gap:10px}
.m{width:22px;height:22px;border-radius:50%;display:grid;place-items:center;color:#fff;font-size:.8rem;font-weight:700;margin-top:1px}
.m.ok{background:var(--ok)}.m.no{background:var(--falta)}
li b{display:block}li span{color:var(--suave);font-size:.9rem}
ol{margin:0;padding-left:20px;display:grid;gap:8px}
.cta{display:flex;flex-wrap:wrap;gap:14px;align-items:center;justify-content:space-between;padding:22px 30px;background:#fafaf8}
.cta b{font-size:1.05rem}.cta a{background:var(--cor);color:#fff;text-decoration:none;padding:11px 18px;border-radius:999px;font-weight:600}
footer{padding:14px 30px;color:var(--suave);font-size:.76rem}
@media (max-width:600px){.topo{grid-template-columns:1fr;justify-items:center;text-align:center}}
@media print{body{background:#fff}.folha{box-shadow:none;margin:0}}
</style></head><body><div class="folha">
<header><small>Diagnóstico de presença digital · __DATA__</small><h1>__NEGOCIO__</h1><p>__LOCAL__</p></header>
<div class="topo"><div class="gauge" style="--v:__NOTA__"><div><div><b>__NOTA__</b><br><span>de 100</span></div></div></div>
<div><h2>__RESUMO__</h2><p>Este é um retrato de como o negócio aparece hoje para quem procura no Google, no Maps e nas redes. Os dados são públicos e foram verificados em __DATA__.</p></div></div>
<section><h3>O que encontramos</h3><ul>__ITENS__</ul></section>
<section><h3>Melhorias de maior impacto</h3><ol>__MELHORIAS__</ol></section>
<div class="cta"><div><b>__VENDEDOR__</b><br><span style="color:var(--suave)">__CHAMADA__</span></div>__BOTAO__</div>
<footer>Diagnóstico gerado com a skill gratuita Scout de negócios sem site, da Trentin Solutions · trentinsolutions.com.br</footer>
</div></body></html>"""


def gerar_diag(r, marca, ofertas, pasta):
    x = r["_r"]
    pres = x["presenca"]
    resumo = ("A presença digital ainda está no começo." if pres < 35 else
              "Há uma boa base, com espaços claros para melhorar." if pres < 70 else
              "A presença digital está bem cuidada.")
    if x["site_status"] == "sem_site" and int(r.get("avaliacoes") or 0) >= 50:
        resumo = f"Muita gente já conhece e avalia o negócio ({r.get('avaliacoes')} avaliações), mas ele ainda não tem um site próprio."
    itens = "".join(f'<li><span class="m {"ok" if ok else "no"}">{"✓" if ok else "!"}</span><div><b>{html.escape(t)}</b><span>{html.escape(exp)}</span></div></li>'
                    for _, t, ok, exp in x["itens"])
    sugestao = {"Site próprio": "Ter uma página própria com fotos, serviços, horário, mapa e botão de WhatsApp.",
                "Site seguro (https)": "Ativar o certificado de segurança (https) para tirar o aviso de 'não seguro'.",
                "Abre bem no celular": "Ajustar o site para celular: texto legível e botões fáceis de tocar.",
                "Site atualizado": "Atualizar fotos, serviços e informações do site.",
                "Botão de WhatsApp no site": "Colocar um botão de WhatsApp visível em todas as páginas.",
                "Perfil no Google Maps": "Criar e verificar o perfil no Google Maps.",
                "Horário de funcionamento": "Cadastrar o horário de funcionamento no Google.",
                "Fotos no perfil (10 ou mais)": "Publicar boas fotos do espaço, da equipe e dos produtos no Google.",
                "Responde às avaliações": "Responder às avaliações, inclusive as positivas.",
                "Instagram ativo": "Manter o Instagram com posts regulares, ao menos uma vez por semana.",
                "WhatsApp divulgado": "Divulgar o WhatsApp no Google e nas redes."}
    mel = "".join(f"<li>{html.escape(sugestao.get(t, t))}</li>" for _, t, _, _ in x["melhorias"]) or "<li>Manter o que já está funcionando bem.</li>"
    wa = re.sub(r"\D", "", marca.get("whatsapp", ""))
    if wa and len(wa) in (10, 11):
        wa = "55" + wa
    msg = quote(f"Oi! Recebi o diagnóstico da {r.get('nome', '')} e quero conversar.")
    botao = f'<a href="https://wa.me/{wa}?text={msg}">Conversar no WhatsApp</a>' if wa else ""
    pagina = (DIAG.replace("__COR__", marca.get("cor", "#3a6ab4")).replace("__NEGOCIO__", html.escape(r.get("nome", "")))
              .replace("__LOCAL__", html.escape(" · ".join(x for x in (r.get("nicho"), r.get("bairro"), r.get("cidade")) if x)))
              .replace("__DATA__", HOJE.strftime("%d/%m/%Y")).replace("__NOTA__", str(pres)).replace("__RESUMO__", html.escape(resumo))
              .replace("__ITENS__", itens).replace("__MELHORIAS__", mel).replace("__VENDEDOR__", html.escape(marca.get("nome", "")))
              .replace("__CHAMADA__", html.escape(marca.get("chamada", "Posso ajudar a resolver esses pontos.")))
              .replace("__BOTAO__", botao))
    nome = os.path.join(pasta, f"diagnostico_{slug(r.get('nome'))}.html")
    with open(nome, "w", encoding="utf-8") as f:
        f.write(pagina)
    return nome


def cmd_gerar(a):
    with open(a.briefing, encoding="utf-8") as f:
        briefing = json.load(f)
    marca = {}
    if a.marca and os.path.exists(a.marca):
        with open(a.marca, encoding="utf-8") as f:
            marca = json.load(f)
    ofertas = briefing.get("ofertas") or ["site"]
    regs = ler(a.entrada)
    nichos_alvo = [norm(n) for n in briefing.get("nichos", [])]
    for r in regs:
        r["_r"] = oportunidade(r, ofertas)
        r["_r"]["melhorias"] = melhorias(r["_r"], ofertas)
    nichos = tabela_nichos(regs)
    lista = [r for r in regs if not nichos_alvo or norm(r.get("nicho")) in nichos_alvo]
    lista.sort(key=lambda r: (-r["_r"]["oportunidade"], -int(r.get("avaliacoes") or 0)))
    if a.limite:
        lista = lista[:a.limite]
    os.makedirs(a.saida, exist_ok=True)
    pasta_d = os.path.join(a.saida, "diagnosticos")
    os.makedirs(pasta_d, exist_ok=True)
    feitos = 0
    for r in lista:
        if feitos < a.diagnosticos and r["_r"]["prioridade"] in ("Quente", "Morno"):
            r["_r"]["diag_arquivo"] = os.path.relpath(gerar_diag(r, marca, ofertas, pasta_d), a.saida)
            feitos += 1
    planilha = os.path.join(a.saida, f"scout_{slug(briefing.get('cidade', 'cidade'))}.xlsx")
    try:
        gravar_xlsx(lista, nichos, planilha, briefing)
    except ImportError:
        print("openpyxl não instalado (pip install openpyxl); planilha não gerada.")
    q = sum(1 for r in lista if r["_r"]["prioridade"] == "Quente")
    sem = sum(1 for r in lista if r["_r"]["site_status"] in ("sem_site", "so_rede_social", "erro"))
    print(f"{len(regs)} negócios avaliados · {len(lista)} na lista · {sem} sem site de verdade · {q} quentes · {feitos} diagnósticos -> {a.saida}")
    for r in lista[:5]:
        x = r["_r"]
        print(f"  {x['oportunidade']:>3} {x['prioridade']:<6} {r.get('nome', '')[:30]:<30} presença {x['presenca']:>3} · {r.get('avaliacoes', 0)} avaliações · falta: {', '.join(i[1] for i in x['melhorias'][:2])}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s1 = sub.add_parser("nichos")
    s1.add_argument("entrada")
    s1.add_argument("--ofertas", default="site")
    s2 = sub.add_parser("gerar")
    s2.add_argument("entrada")
    s2.add_argument("--briefing", required=True)
    s2.add_argument("--marca")
    s2.add_argument("--saida", default="scout")
    s2.add_argument("--limite", type=int, default=0)
    s2.add_argument("--diagnosticos", type=int, default=10)
    a = ap.parse_args()
    {"nichos": cmd_nichos, "gerar": cmd_gerar}[a.cmd](a)


if __name__ == "__main__":
    main()
