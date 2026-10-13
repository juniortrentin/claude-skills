#!/usr/bin/env python3
"""Pontua decisores encontrados e monta a planilha da skill Prospecção via LinkedIn.

Trentin Solutions · criada por Ivaldo Trentin · trentinsolutions.com.br
Uso livre, mantendo este crédito (ver LICENSE.md).

Uso:
  python montar_lista.py trabalho/pessoas.jsonl --briefing trabalho/briefing.json \
      --saida prospeccao_linkedin.xlsx [--limite 50] [--anterior planilha.xlsx]

Nota (0 a 100) = cargo (até 35) + segmento (até 20) + porte (até 15) + região (até 10) + sinais de compra (até 20).
--anterior aceita uma planilha desta skill ou da skill Conexões do LinkedIn com contexto: quem já está lá é pulado.
"""
import argparse
import csv
import json
import os
import re
import sys
import unicodedata
from datetime import date

COL_DEC = ["Nota", "Prioridade", "Nome", "Cargo", "Senioridade", "Empresa", "Segmento", "Nº de colaboradores",
           "Localização", "Sinais de compra", "Motivo", "No cargo desde", "Fonte", "Confirmado no LinkedIn",
           "LinkedIn", "Site da empresa", "Nota de convite sugerida", "Observação"]
COL_EMP = ["Empresa", "Melhor nota", "Decisores na lista", "Segmento", "Nº de colaboradores", "Localização",
           "Sinais de compra", "Site"]
SENIORIDADE = [
    ("C-level", ["ceo", "cfo", "cto", "coo", "cmo", "chro", "cio", "chief", "presidente", "fundador", "fundadora",
                 "founder", "socio", "socia", "owner", "dono", "dona", "proprietari"]),
    ("Diretoria", ["diretor", "diretora", "director", "vp", "vice presidente", "head"]),
    ("Gerência", ["gerente", "manager", "superintendente"]),
    ("Coordenação", ["coordenador", "coordenadora", "supervisor", "supervisora", "lider", "lead"]),
    ("Especialista", ["especialista", "consultor", "consultora", "analista senior", "senior", "sr"]),
    ("Analista", ["analista", "assistente", "auxiliar", "estagiario", "estagiaria", "trainee"]),
]
SINAIS = {"vagas": "vagas abertas na área", "crescimento": "empresa crescendo", "novo_no_cargo": "novo no cargo",
          "post": "post recente sobre o tema", "investimento": "recebeu investimento", "expansao": "expansão ou nova unidade"}


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", s)).strip()


def url_chave(u):
    u = str(u or "").strip().lower().split("?")[0].rstrip("/")
    m = re.search(r"linkedin\.com/in/([^/]+)", u)
    return "in/" + m.group(1) if m else u


def senioridade(cargo):
    c = " " + norm(cargo) + " "
    for nivel, termos in SENIORIDADE:
        if any(f" {t} " in c or c.strip().startswith(t + " ") for t in termos):
            return nivel
    return ""


def faixa(n):
    try:
        n = int(re.sub(r"\D", "", str(n)) or 0)
    except ValueError:
        return ""
    if not n:
        return ""
    return "1-10" if n <= 10 else "11-50" if n <= 50 else "51-200" if n <= 200 else "201-500" if n <= 500 \
        else "501-1000" if n <= 1000 else "1001-5000" if n <= 5000 else "5000+"


def pontuar(p, b):
    motivos, nota = [], 0
    cargo = norm(p.get("cargo"))
    # cargo: lista do briefing com peso (principal 35, secundário 20); sem lista, pela senioridade
    melhor = 0
    for alvo in b.get("cargos", []):
        termo = norm(alvo["termo"] if isinstance(alvo, dict) else alvo)
        peso = (alvo.get("peso", "principal") if isinstance(alvo, dict) else "principal")
        if termo and re.search(rf"\b{re.escape(termo)}\b", cargo):
            melhor = max(melhor, 35 if peso == "principal" else 20)
    niv = senioridade(p.get("cargo"))
    if not b.get("cargos"):
        melhor = {"C-level": 35, "Diretoria": 30, "Gerência": 22, "Coordenação": 15}.get(niv, 5)
    elif not melhor and niv and niv in b.get("senioridades", []):
        melhor = 15
    nota += melhor
    if melhor >= 30:
        motivos.append("cargo principal")
    elif melhor >= 15:
        motivos.append("cargo secundário")
    segs = [norm(s) for s in b.get("segmentos", [])]
    alvo_seg = norm(p.get("segmento")) + " " + norm(p.get("empresa"))
    if segs:
        if any(s and s in alvo_seg for s in segs):
            nota += 20
            motivos.append("segmento-alvo")
    else:
        nota += 10
    portes = b.get("portes", [])
    fp = faixa(p.get("colaboradores"))
    if portes:
        if fp in portes:
            nota += 15
            motivos.append(f"porte {fp}")
    elif fp:
        nota += 7
    regs = [norm(r) for r in b.get("regioes", [])]
    loc = norm(p.get("localizacao"))
    if regs:
        if any(r and re.search(rf"\b{re.escape(r)}\b", loc) for r in regs):
            nota += 10
            motivos.append("região-alvo")
    else:
        nota += 5
    sinais = [s for s in (p.get("sinais") or []) if isinstance(s, dict) and s.get("evidencia")]
    pesados = [s for s in sinais if s.get("tipo") in b.get("sinais", SINAIS.keys())]
    nota += min(20, 10 * len(pesados))
    if pesados:
        motivos.append(" e ".join(SINAIS.get(s.get("tipo"), s.get("tipo", "sinal")) for s in pesados[:2]))
    nota = min(100, nota)
    p.update({"nota": nota, "prioridade": "Alta" if nota >= 70 else "Média" if nota >= 45 else "Baixa",
              "senioridade": niv, "motivo": ", ".join(motivos) or "pouca aderência ao briefing",
              "sinais_txt": "; ".join(f"{SINAIS.get(s.get('tipo'), s.get('tipo'))}: {s['evidencia']}" for s in sinais)})
    return p


def anteriores(caminho):
    if not caminho or not os.path.exists(caminho):
        return set()
    urls = set()
    if caminho.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(caminho, read_only=True, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, str) and "linkedin.com/in/" in v.lower():
                        urls.add(url_chave(v))
    else:
        with open(caminho, encoding="utf-8-sig") as f:
            for row in csv.reader(f):
                for v in row:
                    if "linkedin.com/in/" in v.lower():
                        urls.add(url_chave(v))
    return urls


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("entrada")
    ap.add_argument("--briefing", required=True)
    ap.add_argument("--saida", default="prospeccao_linkedin.xlsx")
    ap.add_argument("--limite", type=int, default=0)
    ap.add_argument("--anterior")
    ap.add_argument("--plano", default="trabalho/plano_buscas.md")
    a = ap.parse_args()
    with open(a.briefing, encoding="utf-8") as f:
        b = json.load(f)
    pessoas, vistos = [], {}
    with open(a.entrada, encoding="utf-8") as f:
        for n, l in enumerate(f, 1):
            l = l.strip()
            if not l:
                continue
            try:
                p = json.loads(l)
            except json.JSONDecodeError as e:
                print(f"Aviso: linha {n} ignorada ({e})", file=sys.stderr)
                continue
            k = url_chave(p.get("url")) or norm(p.get("nome", "") + p.get("empresa", ""))
            if k in vistos:
                vistos[k].update({x: y for x, y in p.items() if y and not vistos[k].get(x)})
            else:
                vistos[k] = p
                pessoas.append(p)
    ja = anteriores(a.anterior)
    pulados = [p for p in pessoas if url_chave(p.get("url")) in ja]
    pessoas = [pontuar(p, b) for p in pessoas if url_chave(p.get("url")) not in ja]
    pessoas.sort(key=lambda p: (-p["nota"], p.get("empresa", "")))
    if a.limite:
        pessoas = pessoas[:a.limite]
    empresas = {}
    for p in pessoas:
        e = empresas.setdefault(norm(p.get("empresa")) or p.get("nome"), {
            "empresa": p.get("empresa", ""), "melhor": 0, "n": 0, "segmento": p.get("segmento", ""),
            "colaboradores": p.get("colaboradores", ""), "localizacao": p.get("localizacao", ""), "sinais": set(),
            "site": p.get("site", "")})
        e["melhor"] = max(e["melhor"], p["nota"])
        e["n"] += 1
        for s in p.get("sinais") or []:
            if isinstance(s, dict) and s.get("tipo") and s.get("evidencia"):
                e["sinais"].add(SINAIS.get(s["tipo"], s["tipo"]))
    plano = []
    if a.plano and os.path.exists(a.plano):
        with open(a.plano, encoding="utf-8") as f:
            plano = [l.rstrip("\n") for l in f]
    gravar(pessoas, sorted(empresas.values(), key=lambda e: -e["melhor"]), plano, b, a.saida)
    altas = sum(1 for p in pessoas if p["prioridade"] == "Alta")
    print(f"{len(pessoas)} decisores · {len(empresas)} empresas · {altas} com prioridade alta · {len(pulados)} já estavam na planilha anterior -> {a.saida}")
    for p in pessoas[:5]:
        print(f"  {p['nota']:>3} {p['prioridade']:<5} {p.get('nome', '')} · {p.get('cargo', '')} · {p.get('empresa', '')} — {p['motivo']}")


def gravar(pessoas, empresas, plano, b, saida):
    linhas = [[p["nota"], p["prioridade"], p.get("nome", ""), p.get("cargo", ""), p["senioridade"], p.get("empresa", ""),
               p.get("segmento", ""), p.get("colaboradores", ""), p.get("localizacao", ""), p["sinais_txt"], p["motivo"],
               p.get("no_cargo_desde", ""), p.get("fonte", ""), "sim" if p.get("confirmado") else "não",
               p.get("url", ""), p.get("site", ""), p.get("nota_convite", ""), p.get("observacao", "")] for p in pessoas]
    if not saida.lower().endswith(".xlsx"):
        with open(saida, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(COL_DEC)
            w.writerows(linhas)
        return
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    azul, branco = PatternFill("solid", fgColor="3A6AB4"), Font(bold=True, color="FFFFFF", name="Arial")
    corpo = Font(name="Arial", size=10)
    cores = {"Alta": "D8F3E6", "Média": "FCEFD6", "Baixa": "ECEFF4"}
    wb = Workbook()
    wb.properties.creator = "Trentin Solutions · Prospecção via LinkedIn"
    wb.properties.description = "Gerada pela skill gratuita Prospecção via LinkedIn, da Trentin Solutions (trentinsolutions.com.br)."
    ws = wb.active
    ws.title = "Decisores"
    ws.append(COL_DEC)
    for l in linhas:
        ws.append(l)
    for i, w_ in enumerate([7, 10, 24, 30, 13, 26, 18, 12, 22, 46, 34, 13, 14, 11, 40, 30, 50, 30], 1):
        ws.column_dimensions[get_column_letter(i)].width = w_
    for c in ws[1]:
        c.fill, c.font = azul, branco
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = corpo
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (10, 11, 17))
        row[1].fill = PatternFill("solid", fgColor=cores.get(row[1].value, "FFFFFF"))
        row[0].font = Font(name="Arial", size=10, bold=True)
        for j in (14, 15):
            if row[j].value and str(row[j].value).startswith("http"):
                row[j].hyperlink = row[j].value
                row[j].font = Font(name="Arial", size=10, color="3A6AB4", underline="single")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    we = wb.create_sheet("Empresas")
    we.append(COL_EMP)
    for e in empresas:
        we.append([e["empresa"], e["melhor"], e["n"], e["segmento"], e["colaboradores"], e["localizacao"],
                   "; ".join(sorted(e["sinais"])), e["site"]])
    for i, w_ in enumerate([30, 11, 10, 20, 12, 22, 44, 30], 1):
        we.column_dimensions[get_column_letter(i)].width = w_
    for c in we[1]:
        c.fill, c.font = azul, branco
    we.freeze_panes = "B2"

    wbf = wb.create_sheet("Briefing")
    wbf.append(["Critério", "Valor"])
    for c in wbf[1]:
        c.fill, c.font = azul, branco
    cargos = ", ".join((c["termo"] + (" (secundário)" if c.get("peso") == "secundario" else "")) if isinstance(c, dict) else c
                       for c in b.get("cargos", []))
    for nome, val in (("Oferta", b.get("oferta", "")), ("Segmentos", ", ".join(b.get("segmentos", []))), ("Cargos", cargos),
                      ("Senioridades", ", ".join(b.get("senioridades", []))), ("Porte", ", ".join(b.get("portes", []))),
                      ("Regiões", ", ".join(b.get("regioes", []))),
                      ("Sinais de compra", ", ".join(SINAIS.get(s, s) for s in b.get("sinais", [])))):
        wbf.append([nome, val or "(não definido)"])
    wbf.append([])
    wbf.append(["Nota", "cargo até 35 + segmento até 20 + porte até 15 + região até 10 + sinais até 20. Alta ≥ 70, Média ≥ 45."])
    wbf.append([])
    wbf.append(["Plano de busca usado", ""])
    for l in plano:
        wbf.append(["", l])
    wbf.column_dimensions["A"].width = 22
    wbf.column_dimensions["B"].width = 100

    wc = wb.create_sheet("Como usar")
    for t in ["Prospecção via LinkedIn · trentinsolutions.com.br", "",
              "Decisores: uma pessoa por linha, da maior para a menor nota. Prioridade Alta = abordar primeiro.",
              "Sinais de compra trazem a evidência encontrada (vaga, post, crescimento, cargo novo). Confira antes de abordar.",
              "Fonte: de onde a pessoa veio (Apollo, Google ou LinkedIn). 'Confirmado no LinkedIn' = perfil aberto e conferido.",
              "Nota de convite sugerida: rascunho com até 300 caracteres. Personalize e envie você mesmo.",
              "Depois que a pessoa aceitar, a skill Conexões do LinkedIn com contexto completa os dados de contato.",
              "LGPD: prospecção B2B por legítimo interesse, só dados profissionais. Remova quem pedir para sair.", "",
              "Skill gratuita da Trentin Solutions, criada por Ivaldo Trentin · trentinsolutions.com.br",
              "Esta planilha te ajudou? Um Pix de qualquer valor em agradecimento é bem-vindo.",
              "Chave Pix aleatória: b73301d4-682f-46a0-857a-3a7653c2befb · Ivaldo Idalgo Trentin Junior Ltda"]:
        wc.append([t])
    wc["A1"].font = Font(name="Arial", size=13, bold=True, color="3A6AB4")
    wc.column_dimensions["A"].width = 115
    wb.save(saida)


if __name__ == "__main__":
    main()
