#!/usr/bin/env python3
"""Preço que dá lucro e estoque que gira: preço certo por produto, raio-X do estoque
parado e plano de queima em degraus.

Trentin Solutions · criada por Ivaldo Trentin · trentinsolutions.com.br
Uso livre, mantendo este crédito (ver LICENSE.md).

Uso:
  python precos.py preco --custo 40 [--frete 2] [--canal credito] [--config trabalho/config.json]
      Conta rápida de um produto: para onde vai cada R$ 100 e os três preços.
  python precos.py calcular trabalho/produtos.csv --config trabalho/config.json --saida resultado/
      Planilha completa (Resumo, Preços, Estoque, Contas, Como usar) e resumo.json.

Formato dos arquivos em references/campos.md.
"""
import argparse
import csv
import json
import math
import os
import sys
from datetime import date, datetime, timedelta

HOJE = date.today()

CONFIG_PADRAO = {
    "imposto": 6.0,            # % sobre o que vende
    "gastos_fixos_mes": 0.0,   # aluguel, luz, internet, salários, pró-labore...
    "vendas_mes": 0.0,         # quanto entra por mês, em média
    "lucro": 20.0,             # % que a pessoa quer que sobre limpo
    "guardar_mes": 2.0,        # % do valor que cada mês parado custa
    "canal_padrao": "credito",
    "canais": {
        "pix": {"pct": 0.0, "fixo": 0.0},
        "dinheiro": {"pct": 0.0, "fixo": 0.0},
        "debito": {"pct": 1.5, "fixo": 0.0},
        "credito": {"pct": 3.5, "fixo": 0.0},
        "marketplace": {"pct": 18.0, "fixo": 0.0},
    },
    "degrau_dias": 15,
}

# meses para vender tudo: até ok = tranquilo; até atencao = devagar; acima = parado
LIMITES = {"duravel": (3, 6), "moda": (2, 4), "validade": (1, 2)}
IDADE_NOVO = 45          # dias: produto recém-chegado ainda não é avaliado como parado
IDADE_PARADO = {"duravel": 180, "moda": 90, "validade": 60}


# ------------------------------------------------------------ leitura
def num(v, padrao=0.0):
    if v is None:
        return padrao
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("R$", "").replace("%", "").strip()
    if not s:
        return padrao
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return padrao


def data(v):
    if not v:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for f in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def ler_config(caminho):
    cfg = json.loads(json.dumps(CONFIG_PADRAO))
    if caminho and os.path.exists(caminho):
        with open(caminho, encoding="utf-8") as f:
            dado = json.load(f)
        canais = dado.pop("canais", None)
        cfg.update(dado)
        if canais:
            cfg["canais"] = {k.lower(): {"pct": num(v.get("pct")), "fixo": num(v.get("fixo"))} if isinstance(v, dict)
                             else {"pct": num(v), "fixo": 0.0} for k, v in canais.items()}
    for k in ("imposto", "gastos_fixos_mes", "vendas_mes", "lucro", "guardar_mes"):
        cfg[k] = num(cfg.get(k))
    return cfg


def fixos_pct(cfg):
    """Quanto de cada venda vai para as contas fixas do mês (%)."""
    if cfg["vendas_mes"] > 0:
        return 100.0 * cfg["gastos_fixos_mes"] / cfg["vendas_mes"]
    return 0.0


def ler_produtos(caminho):
    ext = os.path.splitext(caminho)[1].lower()
    linhas = []
    if ext == ".jsonl":
        with open(caminho, encoding="utf-8") as f:
            linhas = [json.loads(l) for l in f if l.strip()]
    elif ext in (".xlsx", ".xlsm"):
        import openpyxl
        ws = openpyxl.load_workbook(caminho, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        cab = [str(c or "").strip().lower() for c in rows[0]]
        linhas = [dict(zip(cab, r)) for r in rows[1:] if any(r)]
    else:
        with open(caminho, encoding="utf-8-sig") as f:
            amostra = f.read(2048)
            f.seek(0)
            sep = ";" if amostra.count(";") > amostra.count(",") else ","
            linhas = list(csv.DictReader(f, delimiter=sep))
    out = []
    for l in linhas:
        l = {str(k).strip().lower(): v for k, v in l.items() if k}
        if not str(l.get("nome") or "").strip():
            continue
        tipo = str(l.get("tipo") or "duravel").strip().lower()
        tipo = {"durável": "duravel", "perecivel": "validade", "perecível": "validade", "colecao": "moda", "coleção": "moda"}.get(tipo, tipo)
        if tipo not in LIMITES:
            tipo = "duravel"
        out.append({
            "nome": str(l["nome"]).strip(),
            "categoria": str(l.get("categoria") or "").strip(),
            "tipo": tipo,
            "custo": num(l.get("custo")),
            "frete": num(l.get("frete")) + num(l.get("embalagem")),
            "quantidade": num(l.get("quantidade")),
            "data_compra": data(l.get("data_compra")) or data(l.get("ultima_entrada")),
            "vendidos_90d": num(l.get("vendidos_90d")),
            "preco_atual": num(l.get("preco_atual")),
            "canal": str(l.get("canal") or "").strip().lower(),
            "validade": data(l.get("validade")),
        })
    return out


# ------------------------------------------------------------ contas
def ate_90(x):
    """Arredonda para cima até o próximo R$ X,90."""
    if x <= 0:
        return 0.0
    return math.ceil(x - 0.9 - 1e-9) + 0.9


def abaixo_90(x):
    """Arredonda para baixo até o R$ X,90 anterior."""
    return max(0.9, math.floor(x - 0.9 + 1e-9) + 0.9)


def taxa(cfg, canal):
    c = cfg["canais"].get(canal) or cfg["canais"].get(cfg["canal_padrao"]) or {"pct": 0.0, "fixo": 0.0}
    return c["pct"], c["fixo"]


def precos(custo_total, cfg, canal):
    """Três preços: devolve o que pagou, paga as contas, dá o lucro desejado."""
    pct, fixo = taxa(cfg, canal)
    imp, fx, lu = cfg["imposto"], fixos_pct(cfg), cfg["lucro"]
    base = custo_total + fixo
    res = {}
    for nome, soma in (("devolve", imp + pct), ("minimo", imp + pct + fx), ("sugerido", imp + pct + fx + lu)):
        livre = 1 - soma / 100
        res[nome] = base / livre if livre > 0.05 else None
    res["sugerido_90"] = ate_90(res["sugerido"]) if res["sugerido"] else None
    return res


def de_cada_100(cfg, canal):
    pct, _ = taxa(cfg, canal)
    partes = [("imposto", cfg["imposto"]), ("maquininha ou plataforma", pct),
              ("contas fixas do mês", fixos_pct(cfg)), ("seu lucro", cfg["lucro"])]
    resto = 100 - sum(p for _, p in partes)
    return partes, resto


def sobra_por_venda(preco, custo_total, cfg, canal):
    pct, fixo = taxa(cfg, canal)
    return preco * (1 - (cfg["imposto"] + pct + fixos_pct(cfg)) / 100) - custo_total - fixo


def analisar(p, cfg):
    canal = p["canal"] or cfg["canal_padrao"]
    ct = p["custo"] + p["frete"]
    pr = precos(ct, cfg, canal)
    r = {"canal": canal, "custo_total": ct, **pr}
    atual = p["preco_atual"]
    if atual > 0:
        r["sobra"] = sobra_por_venda(atual, ct, cfg, canal)
        r["sobra_pct"] = r["sobra"] / atual
        if pr["devolve"] and atual < pr["devolve"] - 0.005:
            r["preco_situacao"] = "Prejuízo"
            r["preco_texto"] = f"Perde dinheiro a cada venda: abaixo de {brl(pr['devolve'])} não devolve nem o que você pagou."
        elif pr["minimo"] and atual < pr["minimo"] - 0.005:
            r["preco_situacao"] = "Não paga as contas"
            r["preco_texto"] = f"Devolve o que pagou, mas não ajuda a pagar as contas do mês. O mínimo é {brl(pr['minimo'])}."
        elif pr["sugerido"] and atual < pr["sugerido"] - 0.005:
            r["preco_situacao"] = "Lucro baixo"
            r["preco_texto"] = f"Paga as contas, mas sobra menos do que você quer. Sugerido: {brl(pr['sugerido_90'])}."
        else:
            r["preco_situacao"] = "Lucro ok"
            r["preco_texto"] = "O preço cobre tudo e deixa o lucro que você quer."
    else:
        r["sobra"] = r["sobra_pct"] = None
        r["preco_situacao"] = "Sem preço"
        r["preco_texto"] = f"Sugerido: {brl(pr['sugerido_90'])}." if pr["sugerido_90"] else "Confira as porcentagens: somadas, passam de 95%."

    # ------- estoque
    q = p["quantidade"]
    idade = (HOJE - p["data_compra"]).days if p["data_compra"] else None
    por_mes = p["vendidos_90d"] / 3
    meses = q / por_mes if por_mes > 0 else (math.inf if q > 0 else 0)
    parado = ct * q
    r.update(idade=idade, por_mes=por_mes, meses=meses, parado=parado,
             custo_esperar_mes=parado * cfg["guardar_mes"] / 100)
    ok, at = LIMITES[p["tipo"]]
    vence = (p["validade"] - HOJE).days if p["validade"] else None
    if q <= 0:
        faixa = "Sem estoque"
    elif vence is not None and (vence <= 0 or (por_mes > 0 and meses * 30 > vence) or (por_mes == 0 and vence < 60)):
        faixa = "Vence antes de vender"
    elif idade is not None and idade < IDADE_NOVO and meses > ok:
        faixa = "Recém-chegado"
    elif meses <= ok:
        faixa = "Vendendo bem"
    elif idade is not None and idade >= IDADE_PARADO[p["tipo"]]:
        faixa = "Parado"  # está há muito tempo na prateleira e ainda não sai no ritmo bom
    elif meses <= at or por_mes == 0:
        faixa = "Devagar"
    else:
        faixa = "Parado"
    r["faixa"] = faixa
    r["vence_dias"] = vence

    # ------- plano
    degraus = []
    ref = atual if atual > 0 else (pr["sugerido_90"] or 0)
    piso = pr["devolve"] or ct
    if faixa in ("Parado", "Vence antes de vender") and ref > 0:
        if faixa == "Parado":
            # 20% e 35% de desconto, sem passar do piso; o último degrau só devolve o que pagou
            alvos = [ate_90(max(ref * 0.8, piso)), ate_90(max(ref * 0.65, piso)), math.ceil(piso * 10) / 10]
            passo = cfg["degrau_dias"]
        else:
            # vencendo: o último degrau pode ficar abaixo do que pagou (recupera 70%)
            alvos = [ate_90(max(ref * 0.8, piso)), math.ceil(piso * 10) / 10, abaixo_90(ct * 0.7)]
            passo = max(3, min(cfg["degrau_dias"], ((vence or 45) - 1) // 3))
        for v in alvos:
            if v >= ref:
                continue
            if degraus and v >= degraus[-1]["preco"] - 1:
                # muito perto do degrau anterior: fica só o menor
                if v < degraus[-1]["preco"]:
                    degraus[-1]["preco"] = v
                continue
            degraus.append({"preco": v, "ate": HOJE + timedelta(days=passo * (len(degraus) + 1))})
    r["degraus"] = degraus
    pct, fixo = taxa(cfg, canal)
    if degraus:
        ult = degraus[-1]["preco"]
        r["volta"] = q * (ult * (1 - (cfg["imposto"] + pct) / 100) - fixo)
    else:
        r["volta"] = None

    if faixa == "Vendendo bem":
        acao = "Está vendendo bem. Mantenha" + (f" e, se puder, suba para {brl(pr['sugerido_90'])}." if r["preco_situacao"] in ("Prejuízo", "Não paga as contas", "Lucro baixo") else " o preço.")
    elif faixa == "Devagar":
        acao = ("Está devagar. Monte um kit com um produto que vende bem, use como brinde em compras maiores ou dê destaque no catálogo."
                + (f" Se baixar o preço, não passe de {brl(pr['minimo'])}." if pr["minimo"] else ""))
    elif faixa == "Recém-chegado":
        acao = "Chegou há pouco tempo. Espere mais umas semanas antes de mexer no preço."
    elif faixa == "Parado":
        meses_txt = "não vendeu nada nos últimos 3 meses" if por_mes == 0 else f"no ritmo de hoje leva {fmt_meses(meses)} para vender tudo"
        if degraus:
            acao = (f"Parado: {meses_txt}. Vale mais o dinheiro de volta agora. Faça a liquidação em degraus; "
                    f"o último só devolve o que você pagou.")
        else:
            acao = (f"Parado: {meses_txt}, e o preço de hoje já não devolve o que você pagou nessa forma de venda. "
                    f"Antes de baixar mais, tente outra forma de venda com taxa menor, um kit ou vender para outro lojista.")
    elif faixa == "Vence antes de vender":
        acao = "Vai vencer antes de vender tudo. Liquide já; no último degrau vale vender abaixo do que pagou para não perder tudo."
    else:
        acao = "Sem estoque."
    r["acao"] = acao
    return r


def brl(v):
    if v is None:
        return "-"
    return "R$ " + f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def fmt_meses(m):
    if m == math.inf:
        return "sem previsão"
    if m < 1:
        return f"{round(m * 30)} dias"
    if m >= 24:
        return f"{m / 12:.0f} anos"
    return f"{m:.0f} meses" if round(m) != 1 else "1 mês"


# ------------------------------------------------------------ comandos
def cmd_preco(a):
    cfg = ler_config(a.config)
    canal = (a.canal or cfg["canal_padrao"]).lower()
    ct = a.custo + a.frete
    pr = precos(ct, cfg, canal)
    partes, resto = de_cada_100(cfg, canal)
    print(f"Produto que custa {brl(ct)} (com frete e embalagem), vendido no {canal}:\n")
    print("De cada R$ 100 que entram:")
    for n, p in partes:
        print(f"  {brl(p):>10}  {n}")
    print(f"  {brl(resto):>10}  pagam o produto\n")
    if resto <= 5:
        print("Atenção: as porcentagens somadas passam de 95%. Revise as contas.")
        return
    print(f"  Devolve o que você pagou:     {brl(pr['devolve'])}")
    print(f"  Paga as contas (mínimo):      {brl(pr['minimo'])}")
    print(f"  Dá o lucro que você quer:     {brl(pr['sugerido'])}  → sugerido {brl(pr['sugerido_90'])}")


def cmd_calcular(a):
    cfg = ler_config(a.config)
    prods = ler_produtos(a.entrada)
    if not prods:
        sys.exit("Nenhum produto encontrado. Confira a coluna 'nome'.")
    for p in prods:
        p["_r"] = analisar(p, cfg)
    os.makedirs(a.saida, exist_ok=True)
    xl = os.path.join(a.saida, f"precos_e_estoque_{HOJE:%Y-%m-%d}.xlsx")
    gravar_xlsx(prods, cfg, xl)
    res = resumo(prods, cfg)
    with open(os.path.join(a.saida, "resumo.json"), "w", encoding="utf-8") as f:
        json.dump(res, f, ensure_ascii=False, indent=2, default=str)
    print(f"Planilha: {xl}")
    print(json.dumps(res, ensure_ascii=False, indent=2, default=str))


def resumo(prods, cfg):
    R = [p["_r"] for p in prods]
    def soma(f, cond=lambda r: True):
        return round(sum(f(r) for r in R if cond(r)), 2)
    faixas = {}
    for r in R:
        faixas[r["faixa"]] = faixas.get(r["faixa"], 0) + 1
    return {
        "produtos": len(R),
        "dinheiro_no_estoque": soma(lambda r: r["parado"]),
        "dinheiro_parado": soma(lambda r: r["parado"], lambda r: r["faixa"] in ("Parado", "Vence antes de vender")),
        "volta_com_a_queima": soma(lambda r: r["volta"] or 0),
        "faixas": faixas,
        "vendidos_no_prejuizo": [p["nome"] for p in prods if p["_r"]["preco_situacao"] == "Prejuízo"],
        "nao_pagam_as_contas": [p["nome"] for p in prods if p["_r"]["preco_situacao"] == "Não paga as contas"],
        "de_cada_100": {n: v for n, v in de_cada_100(cfg, cfg["canal_padrao"])[0]},
        "contas_fixas_pct": round(fixos_pct(cfg), 1),
    }


# ------------------------------------------------------------ planilha
def gravar_xlsx(prods, cfg, caminho):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.utils import get_column_letter as L

    AZ, CI = "3A6AB4", "3BC1CD"
    wb = openpyxl.Workbook()
    cab = Font(bold=True, color="FFFFFF")
    fcab = PatternFill("solid", fgColor=AZ)
    fin = PatternFill("solid", fgColor="EAF3FB")
    fino = Side(style="thin", color="D9DDE3")
    borda = Border(bottom=fino)
    MO = 'R$ #,##0.00'
    PC = '0.0%'
    verde, amar, verm, cinza = "E3F4EA", "FFF4D6", "FBE3DF", "F1F2F4"

    def cabec(ws, cols, larg):
        for i, (c, w) in enumerate(zip(cols, larg), 1):
            x = ws.cell(row=1, column=i, value=c)
            x.font, x.fill = cab, fcab
            x.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[L(i)].width = w
        ws.row_dimensions[1].height = 34
        ws.freeze_panes = "B2"

    # ---------------- Contas (premissas editáveis)
    wc = wb.active
    wc.title = "Contas"
    wc.column_dimensions["A"].width = 38
    wc.column_dimensions["B"].width = 16
    wc.column_dimensions["C"].width = 16
    wc.column_dimensions["D"].width = 60
    wc["A1"] = "Suas contas (mude os números em azul e a planilha recalcula)"
    wc["A1"].font = Font(bold=True, size=13, color=AZ)
    linhas = [
        ("Imposto sobre o que você vende", cfg["imposto"] / 100, PC, "No Simples, a porcentagem da sua guia. MEI que paga valor fixo: 0% e some a guia nas contas fixas."),
        ("Contas fixas do mês", cfg["gastos_fixos_mes"], MO, "Aluguel, luz, internet, salários, seu pró-labore, sistema..."),
        ("Quanto você vende por mês (média)", cfg["vendas_mes"], MO, "Tudo o que entra em vendas, em média."),
        ("Parte de cada venda que paga as contas", "=IF(B5>0,B4/B5,0)", PC, "Calculado: contas fixas ÷ vendas do mês."),
        ("Lucro que você quer", cfg["lucro"] / 100, PC, "O que sobra limpo depois de pagar tudo."),
        ("Quanto custa guardar, por mês", cfg["guardar_mes"] / 100, PC, "Espaço, risco de estragar ou sair de moda e o dinheiro que fica preso."),
    ]
    for i, (n, v, fmtx, ajuda) in enumerate(linhas, 3):
        wc.cell(row=i, column=1, value=n)
        c = wc.cell(row=i, column=2, value=v)
        c.number_format = fmtx
        if not (isinstance(v, str) and v.startswith("=")):
            c.fill, c.font = fin, Font(color=AZ, bold=True)
        wc.cell(row=i, column=4, value=ajuda).font = Font(color="636874", size=9)
    wc["A11"] = "Formas de venda"
    wc["A11"].font = Font(bold=True, color=AZ)
    for j, t in enumerate(("Forma", "Taxa %", "Taxa fixa por venda"), 1):
        x = wc.cell(row=12, column=j, value=t)
        x.font, x.fill = cab, fcab
    r0 = 13
    for i, (k, v) in enumerate(cfg["canais"].items()):
        wc.cell(row=r0 + i, column=1, value=k)
        a = wc.cell(row=r0 + i, column=2, value=v["pct"] / 100)
        b = wc.cell(row=r0 + i, column=3, value=v["fixo"])
        a.number_format, b.number_format = PC, MO
        for x in (a, b):
            x.fill, x.font = fin, Font(color=AZ, bold=True)
    rN = r0 + len(cfg["canais"]) - 1
    wc.cell(row=rN + 2, column=1, value="Confira as taxas com a sua maquininha e com cada plataforma: elas mudam por plano e por parcela.").font = Font(color="636874", size=9, italic=True)
    IMP, FIX, LUC = "Contas!$B$3", "Contas!$B$6", "Contas!$B$7"
    CAN, PCT, FXO = f"Contas!$A${r0}:$A${rN}", f"Contas!$B${r0}:$B${rN}", f"Contas!$C${r0}:$C${rN}"

    # ---------------- Preços
    wp = wb.create_sheet("Preços", 0)
    cols = ["Produto", "Forma de venda", "Custo do produto", "Frete e embalagem", "Custo total", "Taxa %", "Taxa fixa",
            "Preço de hoje", "Sobra limpa por venda", "Sobra %", "Devolve o que pagou", "Mínimo (paga as contas)",
            "Sugerido (com lucro)", "Situação", "O que fazer"]
    cabec(wp, cols, [30, 13, 12, 12, 12, 9, 10, 12, 13, 9, 13, 14, 14, 18, 60])
    for i, p in enumerate(prods, 2):
        r = p["_r"]
        wp.cell(row=i, column=1, value=p["nome"])
        wp.cell(row=i, column=2, value=r["canal"])
        wp.cell(row=i, column=3, value=p["custo"]).number_format = MO
        wp.cell(row=i, column=4, value=p["frete"]).number_format = MO
        wp.cell(row=i, column=5, value=f"=C{i}+D{i}").number_format = MO
        wp.cell(row=i, column=6, value=f"=IFERROR(INDEX({PCT},MATCH(B{i},{CAN},0)),0)").number_format = PC
        wp.cell(row=i, column=7, value=f"=IFERROR(INDEX({FXO},MATCH(B{i},{CAN},0)),0)").number_format = MO
        wp.cell(row=i, column=8, value=p["preco_atual"] or None).number_format = MO
        wp.cell(row=i, column=9, value=f'=IF(H{i}>0,H{i}*(1-{IMP}-F{i}-{FIX})-E{i}-G{i},"")').number_format = MO
        wp.cell(row=i, column=10, value=f'=IF(H{i}>0,I{i}/H{i},"")').number_format = PC
        wp.cell(row=i, column=11, value=f'=IFERROR((E{i}+G{i})/(1-{IMP}-F{i}),"")').number_format = MO
        wp.cell(row=i, column=12, value=f'=IFERROR((E{i}+G{i})/(1-{IMP}-F{i}-{FIX}),"")').number_format = MO
        wp.cell(row=i, column=13, value=f'=IFERROR(ROUNDUP((E{i}+G{i})/(1-{IMP}-F{i}-{FIX}-{LUC})-0.9,0)+0.9,"")').number_format = MO
        wp.cell(row=i, column=14, value=(f'=IF(H{i}="","Sem preço",IF(H{i}<K{i}-0.005,"Prejuízo",IF(H{i}<L{i}-0.005,"Não paga as contas",'
                                         f'IF(H{i}<M{i}-0.95,"Lucro baixo","Lucro ok"))))'))
        wp.cell(row=i, column=15, value=r["preco_texto"]).alignment = Alignment(wrap_text=True, vertical="top")
        for c in range(1, 16):
            wp.cell(row=i, column=c).border = borda
    n = len(prods) + 1
    for txt, cor in (("Prejuízo", verm), ("Não paga as contas", amar), ("Lucro baixo", amar), ("Lucro ok", verde)):
        wp.conditional_formatting.add(f"N2:N{n}", CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=PatternFill("solid", fgColor=cor)))
    wp.auto_filter.ref = f"A1:O{n}"

    # ---------------- Estoque
    we = wb.create_sheet("Estoque", 1)
    cols = ["Produto", "Tipo", "Quantidade", "Comprado em", "Dias no estoque", "Vendidos (3 meses)", "Vende por mês",
            "Tempo para vender tudo", "Dinheiro no estoque", "Custo de esperar (mês)", "Situação",
            "Degrau 1", "Até", "Degrau 2", "Até", "Degrau 3", "Até", "Volta para o caixa", "O que fazer"]
    cabec(we, cols, [30, 10, 10, 12, 10, 11, 10, 14, 14, 13, 18, 11, 11, 11, 11, 11, 11, 14, 70])
    TIPO = {"duravel": "durável", "moda": "moda/coleção", "validade": "com validade"}
    for i, p in enumerate(prods, 2):
        r = p["_r"]
        we.cell(row=i, column=1, value=p["nome"])
        we.cell(row=i, column=2, value=TIPO[p["tipo"]])
        we.cell(row=i, column=3, value=p["quantidade"])
        c = we.cell(row=i, column=4, value=p["data_compra"])
        c.number_format = "dd/mm/yyyy"
        we.cell(row=i, column=5, value=f'=IF(D{i}="","",TODAY()-D{i})')
        we.cell(row=i, column=6, value=p["vendidos_90d"])
        we.cell(row=i, column=7, value=f"=F{i}/3").number_format = "0.0"
        we.cell(row=i, column=8, value=fmt_meses(r["meses"]) if p["quantidade"] > 0 else "-")
        we.cell(row=i, column=9, value=f"=C{i}*Preços!E{i}").number_format = MO
        we.cell(row=i, column=10, value=f"=I{i}*Contas!$B$8").number_format = MO
        we.cell(row=i, column=11, value=r["faixa"])
        for k, d in enumerate(r["degraus"][:3]):
            we.cell(row=i, column=12 + 2 * k, value=d["preco"]).number_format = MO
            x = we.cell(row=i, column=13 + 2 * k, value=d["ate"])
            x.number_format = "dd/mm/yyyy"
        if r["volta"] is not None:
            we.cell(row=i, column=18, value=round(r["volta"], 2)).number_format = MO
        we.cell(row=i, column=19, value=r["acao"]).alignment = Alignment(wrap_text=True, vertical="top")
        for c in range(1, 20):
            we.cell(row=i, column=c).border = borda
    for txt, cor in (("Parado", verm), ("Vence antes de vender", verm), ("Devagar", amar), ("Vendendo bem", verde), ("Recém-chegado", cinza)):
        we.conditional_formatting.add(f"K2:K{n}", CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=PatternFill("solid", fgColor=cor)))
    we.auto_filter.ref = f"A1:S{n}"

    # ---------------- Resumo
    ws = wb.create_sheet("Resumo", 0)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 60
    ws["A1"] = "Preço e estoque · resumo"
    ws["A1"].font = Font(bold=True, size=15, color=AZ)
    ws["A2"] = f"Gerado em {HOJE:%d/%m/%Y}"
    ws["A2"].font = Font(color="636874")
    itens = [
        ("Dinheiro no estoque", f"=SUM(Estoque!I2:I{n})", MO, "Quanto você pagou por tudo que está na prateleira."),
        ("Dinheiro parado (produtos parados ou vencendo)", f'=SUMIF(Estoque!K2:K{n},"Parado",Estoque!I2:I{n})+SUMIF(Estoque!K2:K{n},"Vence antes de vender",Estoque!I2:I{n})', MO, "Dinheiro preso em produto que não está saindo."),
        ("Quanto custa esperar, por mês", f'=SUMIF(Estoque!K2:K{n},"Parado",Estoque!J2:J{n})+SUMIF(Estoque!K2:K{n},"Vence antes de vender",Estoque!J2:J{n})', MO, "O que esse estoque parado custa a cada mês que fica lá."),
        ("Volta para o caixa com a liquidação", f"=SUM(Estoque!R2:R{n})", MO, "Se tudo sair no último degrau (o pior caso)."),
        ("Produtos vendidos no prejuízo", f'=COUNTIF(Preços!N2:N{n},"Prejuízo")', "0", "Cada venda tira dinheiro do seu bolso."),
        ("Produtos que não pagam as contas", f'=COUNTIF(Preços!N2:N{n},"Não paga as contas")', "0", "Devolvem o que você pagou, mas não ajudam no aluguel e nas contas."),
        ("Produtos com lucro ok", f'=COUNTIF(Preços!N2:N{n},"Lucro ok")', "0", ""),
    ]
    for i, (a_, f_, fm, aj) in enumerate(itens, 4):
        ws.cell(row=i, column=1, value=a_).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=f_)
        c.number_format, c.font = fm, Font(bold=True, size=12, color=AZ)
        ws.cell(row=i, column=3, value=aj).font = Font(color="636874")
    ws["A13"] = f"De cada R$ 100 que entram (vendendo no {cfg['canal_padrao']})"
    ws["A13"].font = Font(bold=True, color=AZ)
    pc_c = f"IFERROR(INDEX({PCT},MATCH(\"{cfg['canal_padrao']}\",{CAN},0)),0)"
    partes = [("Imposto", f"=100*{IMP}"), ("Maquininha ou plataforma", f"=100*{pc_c}"), ("Contas fixas do mês", f"=100*{FIX}"),
              ("Seu lucro", f"=100*{LUC}"), ("Sobra para pagar o produto", "=100-SUM(B14:B17)")]
    for i, (a_, f_) in enumerate(partes, 14):
        ws.cell(row=i, column=1, value=a_)
        ws.cell(row=i, column=2, value=f_).number_format = MO
    ws["A20"] = "Feito com a skill Preço que dá lucro e estoque que gira, da Trentin Solutions · trentinsolutions.com.br"
    ws["A20"].font = Font(color="9AA0AA", size=9, italic=True)

    # ---------------- Como usar
    wu = wb.create_sheet("Como usar")
    wu.column_dimensions["A"].width = 110
    textos = [
        ("Como usar esta planilha", True),
        ("1. Na aba Contas, confira os números em azul: imposto, contas do mês, vendas do mês, lucro e as taxas de cada forma de venda. Tudo recalcula sozinho.", False),
        ("2. Na aba Preços, veja a coluna Situação. Comece pelos produtos em vermelho (prejuízo) e amarelo (não pagam as contas).", False),
        ("3. Os três preços: 'Devolve o que pagou' é o piso de uma liquidação; 'Mínimo' paga o produto e a sua parte das contas; 'Sugerido' já tem o seu lucro.", False),
        ("4. Na aba Estoque, os produtos 'Parado' têm uma liquidação em até 3 degraus, de 15 em 15 dias. Se vender antes, ótimo: pare no degrau em que estiver.", False),
        ("5. 'Devagar' ainda não pede desconto: tente kit com um produto que vende bem, brinde em compras maiores ou destaque no catálogo.", False),
        ("6. Rode de novo todo mês com as vendas atualizadas para ver quem saiu da lista.", False),
        ("", False),
        ("Por que vender a preço de custo pode valer a pena?", True),
        ("O dinheiro que você pagou no produto parado já saiu do caixa. Esperar não traz esse dinheiro de volta: só deixa ele preso. Vendendo agora, mesmo sem lucro, o dinheiro volta e pode virar produto que vende rápido e dá lucro de novo.", False),
        ("", False),
        ("Skill gratuita da Trentin Solutions · criada por Ivaldo Trentin · trentinsolutions.com.br · licença CC BY 4.0", False),
    ]
    for i, (t, b) in enumerate(textos, 1):
        c = wu.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True)
        c.font = Font(bold=True, size=12, color=AZ) if b else Font()
    wb.save(caminho)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("preco")
    a.add_argument("--custo", type=float, required=True)
    a.add_argument("--frete", type=float, default=0.0)
    a.add_argument("--canal")
    a.add_argument("--config")
    a.set_defaults(f=cmd_preco)
    b = sub.add_parser("calcular")
    b.add_argument("entrada")
    b.add_argument("--config")
    b.add_argument("--saida", default="resultado")
    b.set_defaults(f=cmd_calcular)
    args = ap.parse_args()
    args.f(args)


if __name__ == "__main__":
    main()
