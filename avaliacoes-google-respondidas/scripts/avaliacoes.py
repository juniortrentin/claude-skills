#!/usr/bin/env python3
"""Avaliações do Google respondidas: confere as respostas sugeridas, organiza os temas
e gera a planilha e o relatório de 1 página.

Trentin Solutions · criada por Ivaldo Trentin · trentinsolutions.com.br
Uso livre, mantendo este crédito (ver LICENSE.md).

Uso:
  python avaliacoes.py conferir trabalho/avaliacoes.jsonl
      Aponta respostas repetidas, longas ou curtas demais, com oferta de desconto,
      dados pessoais, sem assinatura ou faltando.
  python avaliacoes.py gerar trabalho/avaliacoes.jsonl --negocio trabalho/negocio.json --saida resultado/
      Planilha (Responder, Todas, Temas, Como usar) e relatorio.html.

Formato dos arquivos em references/campos.md.
"""
import argparse
import html
import json
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from difflib import SequenceMatcher

HOJE = date.today()

# reserva: se o Claude não marcou os temas, estas palavras ajudam a marcar
TEMAS = {
    "atendimento": ["atendimento", "atendente", "atenciosa", "atencioso", "educad", "simpatic", "grosso", "grossa", "mal atendid", "receptiv", "cordial"],
    "demora": ["demora", "demorou", "esper", "atras", "fila", "lento", "lenta", "rapidez", "rapido", "rapida"],
    "preço": ["preço", "preco", "caro", "cara ", "barato", "custo", "valor", "justo"],
    "qualidade": ["qualidade", "sabor", "gostos", "delicios", "bom demais", "otimo produto", "frio", "fria", "queimad", "estragad", "ruim"],
    "limpeza": ["limp", "sujo", "suja", "higien", "banheiro"],
    "ambiente": ["ambiente", "lugar", "aconchegante", "decora", "musica", "barulh", "espaço", "espaco"],
    "entrega": ["entrega", "delivery", "motoboy", "ifood", "chegou"],
    "estacionamento": ["estacionamento", "estacionar", "vaga"],
    "agendamento": ["agend", "horario marcado", "marcar horario", "remarc"],
    "pós-venda": ["troca", "devolu", "garantia", "retorno", "resolveram", "nao resolveram"],
}
PROIBIDO = [
    (r"\b(desconto|cupom|brinde|cortesia|voucher|vale[- ]compra)\b", "oferece vantagem em público (o Google não permite trocar avaliação por benefício e outros clientes vão pedir o mesmo)"),
    (r"\(?\d{2}\)?\s?9?\d{4}[- ]?\d{4}", "tem número de telefone: prefira 'nos chame no WhatsApp do perfil'"),
    (r"\b\d{3}\.?\d{3}\.?\d{3}-?\d{2}\b", "parece ter CPF"),
    (r"\bpedido\s*(n[ºo°.]?|#)\s*\d+", "cita número de pedido do cliente"),
    (r"\b(mentira|mentiroso|mentirosa|absurdo|ridícul|ridicul)", "tom de briga"),
    (r"\b(processo|advogad|justiça|justica|procon)\b", "fala de processo ou Procon: este caso pede cuidado, não resposta pronta"),
]


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s).strip()


def data(v):
    if not v:
        return None
    s = str(v)[:10]
    for f in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def ler(caminho):
    regs, vistos = [], set()
    with open(caminho, encoding="utf-8") as f:
        for i, l in enumerate(f, 1):
            if not l.strip():
                continue
            try:
                r = json.loads(l)
            except json.JSONDecodeError as e:
                sys.exit(f"Linha {i} com erro de formato: {e}")
            chave = r.get("id") or (norm(r.get("autor")), str(r.get("data")), norm(r.get("texto"))[:40])
            if str(chave) in vistos:
                continue
            vistos.add(str(chave))
            r["estrelas"] = int(r.get("estrelas") or 0)
            r["_data"] = data(r.get("data"))
            if not r.get("temas"):
                t = norm(r.get("texto"))
                r["temas"] = [k for k, ps in TEMAS.items() if any(norm(p) in t for p in ps)]
            if not r.get("sentimento"):
                e = r["estrelas"]
                r["sentimento"] = "positivo" if e >= 4 else "misto" if e == 3 else "negativo"
            # temas podem vir como lista ou como {"tema": "+" ou "-"} (elogio ou reclamação)
            t = r.get("temas") or []
            if isinstance(t, dict):
                r["_pol"] = {k: ("-" if str(v).strip() in ("-", "neg", "negativo", "reclamação", "reclamacao") else "+") for k, v in t.items()}
            else:
                padrao = "+" if r["sentimento"] == "positivo" or (r["sentimento"] == "misto" and r["estrelas"] >= 4) else "-"
                r["_pol"] = {k: padrao for k in t}
            r["temas"] = list(r["_pol"])
            regs.append(r)
    return regs


def precisa_responder(r):
    return not r.get("respondida") and not (r.get("resposta_existente") or "").strip()


def prioridade(r):
    if r.get("cuidado"):
        return "1 · Cuidado"
    if r["estrelas"] <= 2:
        return "2 · Reclamação"
    if r["estrelas"] == 3 or r.get("sentimento") == "misto" or "-" in r.get("_pol", {}).values():
        return "3 · Elogio com ressalva"
    return "4 · Elogio"


# ------------------------------------------------------------ conferir
def problemas(regs, negocio=None):
    out = []
    pend = [r for r in regs if precisa_responder(r)]
    assin = norm((negocio or {}).get("assinatura", ""))
    for r in pend:
        resp = (r.get("resposta") or "").strip()
        nome = r.get("autor") or r.get("id")
        if r.get("cuidado"):
            if resp and not r.get("resposta_aprovada_cuidado"):
                out.append((nome, "caso de cuidado com resposta pronta: revise com o dono antes, frase por frase"))
            continue
        if not resp:
            out.append((nome, "sem resposta sugerida"))
            continue
        palavras = len(resp.split())
        if palavras > 90:
            out.append((nome, f"resposta longa ({palavras} palavras); o ideal é até 60"))
        if r["estrelas"] <= 3 and palavras < 15:
            out.append((nome, "resposta curta demais para uma crítica"))
        for padrao, motivo in PROIBIDO:
            if re.search(padrao, resp, re.I):
                out.append((nome, motivo))
        if assin and assin not in norm(resp):
            out.append((nome, f"falta a assinatura ({negocio['assinatura']})"))
        if resp.isupper():
            out.append((nome, "resposta toda em maiúsculas"))
    # repetidas
    for i, a in enumerate(pend):
        for b in pend[i + 1:]:
            ra, rb = norm(a.get("resposta")), norm(b.get("resposta"))
            if ra and rb and SequenceMatcher(None, ra, rb).ratio() > 0.75:
                out.append((a.get("autor"), f"resposta quase igual à de {b.get('autor')}: varie, cite um detalhe da avaliação"))
    return out


def cmd_conferir(a):
    regs = ler(a.entrada)
    neg = json.load(open(a.negocio, encoding="utf-8")) if a.negocio and os.path.exists(a.negocio) else None
    pend = [r for r in regs if precisa_responder(r)]
    print(f"{len(regs)} avaliações · {len(pend)} sem resposta · {sum(1 for r in pend if r.get('cuidado'))} pedem cuidado")
    ps = problemas(regs, neg)
    if not ps:
        print("Tudo certo: nenhuma resposta com problema.")
    for n, m in ps:
        print(f"- {n}: {m}")
    return 1 if ps else 0


# ------------------------------------------------------------ análise
def analisar(regs):
    tot = len(regs)
    media = sum(r["estrelas"] for r in regs) / tot if tot else 0
    dist = Counter(r["estrelas"] for r in regs)
    resp_antes = sum(1 for r in regs if not precisa_responder(r))
    pend = [r for r in regs if precisa_responder(r)]
    elog, recl = Counter(), Counter()
    for r in regs:
        for t, pol in r["_pol"].items():
            (elog if pol == "+" else recl)[t] += 1
    equipe = Counter()
    for r in regs:
        for p in r.get("equipe") or []:
            if r["sentimento"] != "negativo":
                equipe[p] += 1
    meses = defaultdict(list)
    for r in regs:
        if r["_data"]:
            meses[r["_data"].strftime("%Y-%m")].append(r["estrelas"])
    meses = {m: (round(sum(v) / len(v), 2), len(v)) for m, v in sorted(meses.items())}
    destaques = [r for r in regs if r["estrelas"] == 5 and r.get("destaque")]
    return {"total": tot, "media": round(media, 2), "dist": {k: dist.get(k, 0) for k in (5, 4, 3, 2, 1)},
            "respondidas_antes": resp_antes, "pendentes": len(pend),
            "cuidado": sum(1 for r in pend if r.get("cuidado")),
            "elogios": elog.most_common(), "reclamacoes": recl.most_common(), "equipe": equipe.most_common(5),
            "meses": meses, "destaques": [(r.get("destaque"), r.get("autor")) for r in destaques][:4]}


# ------------------------------------------------------------ planilha
def gravar_xlsx(regs, an, neg, caminho):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.utils import get_column_letter as L
    AZ = "3A6AB4"
    wb = openpyxl.Workbook()
    cab, fcab = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=AZ)
    borda = Border(bottom=Side(style="thin", color="D9DDE3"))
    topo = Alignment(wrap_text=True, vertical="top")

    def aba(ws, cols, larg):
        for i, (c, w) in enumerate(zip(cols, larg), 1):
            x = ws.cell(row=1, column=i, value=c)
            x.font, x.fill = cab, fcab
            x.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[L(i)].width = w
        ws.freeze_panes = "A2"

    # Responder
    ws = wb.active
    ws.title = "Responder"
    cols = ["Prioridade", "Data", "Quem", "Estrelas", "Avaliação", "Temas", "Resposta sugerida", "Por que pede cuidado", "Aprovada?", "Publicada em"]
    aba(ws, cols, [20, 11, 16, 9, 60, 20, 60, 34, 11, 13])
    pend = sorted([r for r in regs if precisa_responder(r)], key=lambda r: (prioridade(r), r["_data"] or date.min))
    for i, r in enumerate(pend, 2):
        vals = [prioridade(r), r["_data"], r.get("autor"), "★" * r["estrelas"], r.get("texto"), ", ".join(r.get("temas") or []),
                r.get("resposta") if not r.get("cuidado") else (r.get("resposta") or "Revisar com o dono antes de responder."),
                r.get("cuidado") or "", "", None]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment, c.border = topo, borda
        ws.cell(row=i, column=2).number_format = "dd/mm/yyyy"
        ws.cell(row=i, column=10).number_format = "dd/mm/yyyy"
    n = len(pend) + 1
    for txt, cor in (("1 · Cuidado", "FBE3DF"), ("2 · Reclamação", "FFF4D6"), ("3 · Elogio com ressalva", "F1F2F4"), ("4 · Elogio", "E3F4EA")):
        ws.conditional_formatting.add(f"A2:A{max(n, 2)}", CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=PatternFill("solid", fgColor=cor)))
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='"Sim,Ajustar,Não responder"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"I2:I{max(n, 2)}")

    # Todas
    wt = wb.create_sheet("Todas")
    aba(wt, ["Data", "Quem", "Estrelas", "Sentimento", "Temas", "Equipe citada", "Avaliação", "Já respondida?"], [11, 16, 9, 11, 22, 16, 70, 12])
    for i, r in enumerate(sorted(regs, key=lambda r: r["_data"] or date.min, reverse=True), 2):
        vals = [r["_data"], r.get("autor"), r["estrelas"], r["sentimento"], ", ".join(r.get("temas") or []),
                ", ".join(r.get("equipe") or []), r.get("texto"), "não" if precisa_responder(r) else "sim"]
        for j, v in enumerate(vals, 1):
            c = wt.cell(row=i, column=j, value=v)
            c.alignment, c.border = topo, borda
        wt.cell(row=i, column=1).number_format = "dd/mm/yyyy"
    m = len(regs) + 1
    wt.auto_filter.ref = f"A1:H{m}"

    # Temas
    wm = wb.create_sheet("Temas")
    aba(wm, ["Tema", "Elogios", "Reclamações", "Saldo"], [22, 11, 13, 10])
    temas = sorted(set([t for t, _ in an["elogios"]] + [t for t, _ in an["reclamacoes"]]))
    E, R = dict(an["elogios"]), dict(an["reclamacoes"])
    for i, t in enumerate(temas, 2):
        wm.cell(row=i, column=1, value=t)
        wm.cell(row=i, column=2, value=E.get(t, 0))
        wm.cell(row=i, column=3, value=R.get(t, 0))
        wm.cell(row=i, column=4, value=f"=B{i}-C{i}")
    k = len(temas) + 3
    wm.cell(row=k, column=1, value="Nota por mês").font = Font(bold=True, color=AZ)
    wm.cell(row=k + 1, column=1, value="Mês").font = Font(bold=True)
    wm.cell(row=k + 1, column=2, value="Nota média").font = Font(bold=True)
    wm.cell(row=k + 1, column=3, value="Avaliações").font = Font(bold=True)
    for j, (mes, (md, q)) in enumerate(an["meses"].items(), k + 2):
        wm.cell(row=j, column=1, value=mes)
        wm.cell(row=j, column=2, value=md).number_format = "0.0"
        wm.cell(row=j, column=3, value=q)
    t0 = j + 2 if an["meses"] else k + 3
    wm.cell(row=t0, column=1, value="Resumo").font = Font(bold=True, color=AZ)
    for jj, (a_, f_) in enumerate((("Nota média", "=AVERAGE(Todas!C2:C%d)" % m), ("Avaliações", "=COUNTA(Todas!B2:B%d)" % m),
                                   ("Sem resposta", '=COUNTIF(Todas!H2:H%d,"não")' % m),
                                   ("Respondidas (%)", '=IFERROR(COUNTIF(Todas!H2:H%d,"sim")/COUNTA(Todas!B2:B%d),0)' % (m, m))), t0 + 1):
        wm.cell(row=jj, column=1, value=a_)
        c = wm.cell(row=jj, column=2, value=f_)
        c.number_format = "0.0%" if "%" in a_ else "0.0" if "Nota" in a_ else "0"

    # Como usar
    wu = wb.create_sheet("Como usar")
    wu.column_dimensions["A"].width = 110
    for i, (t, b) in enumerate([
        ("Como usar esta planilha", True),
        ("1. Na aba Responder, comece pelas linhas vermelhas (pedem cuidado) e amarelas (reclamações). Responda reclamações em até 2 dias.", False),
        ("2. Leia cada resposta, ajuste o que quiser e marque 'Sim' em Aprovada?. Depois copie para o Google (ou peça ao Claude para publicar as aprovadas).", False),
        ("3. 'Pede cuidado' não tem resposta pronta: leve a conversa para o privado e, se for avaliação falsa ou ofensiva, use 'Denunciar avaliação' no perfil.", False),
        ("4. A aba Temas mostra o que mais elogiam e o que mais reclamam. O tema com saldo negativo é o primeiro a melhorar.", False),
        ("5. Rode de novo todo mês para acompanhar a nota e responder as novas.", False),
        ("", False),
        (f"Gerado em {HOJE:%d/%m/%Y} com a skill gratuita Avaliações do Google respondidas, da Trentin Solutions · trentinsolutions.com.br", False),
    ], 1):
        c = wu.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True)
        c.font = Font(bold=True, size=12, color=AZ) if b else Font(color="636874" if i == 8 else "000000")
    wb.save(caminho)


# ------------------------------------------------------------ relatório
REL = r"""<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Avaliações · __NEGOCIO__</title>
<style>
:root{--cor:__COR__;--t:#16181d;--s:#636874;--l:#e6e7ea;--ok:#1f8a5b;--no:#c2412d}
*{box-sizing:border-box}body{margin:0;background:#f4f4f2;color:var(--t);font:15px/1.5 system-ui,-apple-system,"Segoe UI",Roboto,sans-serif}
.f{max-width:820px;margin:24px auto;background:#fff;border-radius:18px;overflow:hidden;box-shadow:0 10px 40px -20px rgba(0,0,0,.25)}
header{background:var(--cor);color:#fff;padding:24px 30px}header small{opacity:.85;text-transform:uppercase;letter-spacing:.06em;font-size:.72rem}
h1{margin:6px 0 0;font-size:1.6rem}
.k{display:grid;grid-template-columns:repeat(4,1fr);border-bottom:1px solid var(--l)}
.k div{padding:18px 22px;border-right:1px solid var(--l)}.k div:last-child{border:0}
.k b{display:block;font-size:1.7rem;line-height:1.1}.k span{color:var(--s);font-size:.8rem}
section{padding:20px 30px;border-bottom:1px solid var(--l)}
h3{margin:0 0 12px;font-size:.76rem;letter-spacing:.08em;text-transform:uppercase;color:var(--s)}
.two{display:grid;grid-template-columns:1fr 1fr;gap:26px}
.bar{display:grid;grid-template-columns:110px 1fr 28px;gap:10px;align-items:center;margin:6px 0;font-size:.9rem}
.bar i{display:block;height:10px;border-radius:99px;background:var(--ok)}.bar.n i{background:var(--no)}
.bar em{font-style:normal;color:var(--s);text-align:right}
.st{display:grid;grid-template-columns:30px 1fr 30px;gap:10px;align-items:center;font-size:.86rem;margin:4px 0}
.st i{display:block;height:8px;border-radius:99px;background:#f2b33d}
ol{margin:0;padding-left:20px;display:grid;gap:8px}
blockquote{margin:0 0 10px;padding:10px 14px;border-left:3px solid var(--cor);background:#fafaf8;border-radius:0 10px 10px 0;font-size:.92rem}
blockquote span{display:block;color:var(--s);font-size:.78rem;margin-top:2px}
.m{display:flex;gap:10px;align-items:flex-end;height:110px}.m div{flex:1;display:grid;gap:4px;justify-items:center;font-size:.72rem;color:var(--s)}
.m i{display:block;width:100%;max-width:42px;background:var(--cor);border-radius:6px 6px 0 0}
footer{padding:14px 30px;color:var(--s);font-size:.74rem}
@media (max-width:640px){.k{grid-template-columns:1fr 1fr}.two{grid-template-columns:1fr}}
@media print{body{background:#fff}.f{box-shadow:none;margin:0}}
</style></head><body><div class="f">
<header><small>Relatório de avaliações do Google · __DATA__</small><h1>__NEGOCIO__</h1></header>
<div class="k">__KPIS__</div>
<section class="two"><div><h3>O que mais elogiam</h3>__ELOG__</div><div><h3>O que mais reclamam</h3>__RECL__</div></section>
<section class="two"><div><h3>Estrelas</h3>__DIST__</div><div><h3>Nota por mês</h3><div class="m">__MESES__</div></div></section>
<section><h3>Melhorias de maior impacto</h3><ol>__MELH__</ol></section>
__EXTRA__
<footer>Relatório gerado com a skill gratuita Avaliações do Google respondidas, da Trentin Solutions · trentinsolutions.com.br</footer>
</div></body></html>"""


def relatorio(an, neg, melhorias, caminho):
    esc = html.escape
    def barras(lista, cls=""):
        if not lista:
            return '<p style="color:var(--s);margin:0">Nada que se repita.</p>'
        mx = max(v for _, v in lista)
        return "".join(f'<div class="bar {cls}"><span>{esc(t.capitalize())}</span><i style="width:{100 * v / mx:.0f}%"></i><em>{v}</em></div>' for t, v in lista[:5])
    tot = an["total"] or 1
    kpis = "".join(f"<div><b>{v}</b><span>{l}</span></div>" for v, l in (
        (f"{an['media']:.1f}".replace(".", ","), "nota média"), (an["total"], "avaliações analisadas"),
        (f"{round(100 * an['respondidas_antes'] / tot)}%", "já tinham resposta"), (an["pendentes"], "respostas prontas para aprovar")))
    mx = max(an["dist"].values() or [1]) or 1
    dist = "".join(f'<div class="st"><span>{k}★</span><i style="width:{100 * v / mx:.0f}%"></i><em style="font-style:normal;color:var(--s)">{v}</em></div>' for k, v in an["dist"].items())
    ms = list(an["meses"].items())[-6:]
    meses = "".join(f'<div><b style="color:var(--t)">{md:.1f}</b><i style="height:{max(6, (md - 1) / 4 * 80):.0f}px"></i>{m[5:]}/{m[2:4]}</div>'.replace(".", ",", 1) for m, (md, q) in ms) or '<p style="color:var(--s)">Sem datas.</p>'
    melh = "".join(f"<li>{esc(x)}</li>" for x in melhorias) or "<li>Manter o que já funciona e responder as novas avaliações.</li>"
    extra = ""
    if an["destaques"]:
        extra += "<section><h3>Frases de clientes</h3>" + "".join(f"<blockquote>“{esc(t)}”<span>{esc(a or '')}</span></blockquote>" for t, a in an["destaques"]) + "</section>"
    if an["equipe"]:
        extra += "<section><h3>Equipe mais elogiada</h3><p style='margin:0'>" + " · ".join(f"<b>{esc(p)}</b> ({v})" for p, v in an["equipe"]) + "</p></section>"
    pg = (REL.replace("__COR__", neg.get("cor", "#3a6ab4")).replace("__NEGOCIO__", esc(neg.get("nome", "Seu negócio")))
          .replace("__DATA__", HOJE.strftime("%d/%m/%Y")).replace("__KPIS__", kpis).replace("__ELOG__", barras(an["elogios"]))
          .replace("__RECL__", barras(an["reclamacoes"], "n")).replace("__DIST__", dist).replace("__MESES__", meses)
          .replace("__MELH__", melh).replace("__EXTRA__", extra))
    with open(caminho, "w", encoding="utf-8") as f:
        f.write(pg)


def cmd_gerar(a):
    regs = ler(a.entrada)
    neg = json.load(open(a.negocio, encoding="utf-8")) if a.negocio and os.path.exists(a.negocio) else {}
    an = analisar(regs)
    melhorias = neg.get("melhorias") or [f"Olhar com atenção: {t} ({v} reclamações)." for t, v in an["reclamacoes"][:3]]
    os.makedirs(a.saida, exist_ok=True)
    xl = os.path.join(a.saida, f"avaliacoes_{HOJE:%Y-%m-%d}.xlsx")
    gravar_xlsx(regs, an, neg, xl)
    relatorio(an, neg, melhorias, os.path.join(a.saida, "relatorio.html"))
    ps = problemas(regs, neg)
    print(f"Planilha: {xl}\nRelatório: {os.path.join(a.saida, 'relatorio.html')}")
    print(json.dumps({k: v for k, v in an.items() if k != "destaques"}, ensure_ascii=False, indent=1))
    if ps:
        print("\nAntes de entregar, ajuste:")
        for n, m in ps:
            print(f"- {n}: {m}")
    return 0


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("conferir")
    c.add_argument("entrada")
    c.add_argument("--negocio")
    c.set_defaults(f=cmd_conferir)
    g = sub.add_parser("gerar")
    g.add_argument("entrada")
    g.add_argument("--negocio")
    g.add_argument("--saida", default="resultado")
    g.set_defaults(f=cmd_gerar)
    a = ap.parse_args()
    sys.exit(a.f(a) or 0)


if __name__ == "__main__":
    main()
