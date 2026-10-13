#!/usr/bin/env python3
"""Monta a planilha de leads da skill Prospecção via Instagram.

Trentin Solutions · criada por Ivaldo Trentin · trentinsolutions.com.br
Uso livre, mantendo este crédito (ver LICENSE.md).

Uso:
  python montar_planilha.py trabalho/perfis.jsonl --saida leads_instagram.xlsx \
      [--limite 50] [--anterior planilha_anterior.xlsx] [--plano trabalho/plano_buscas.md]

Lê uma linha JSON por perfil (formato em references/campos.md), remove duplicados,
pula quem já estava na planilha anterior, calcula a pontuação, ordena e grava
as abas Leads, Pessoas, Buscas e Como usar. Sem openpyxl, grava CSV.
"""
import argparse
import csv
import json
import os
import sys
from datetime import date, datetime

COLUNAS = ["Pontuação", "@", "Nome", "Cidade", "Seguidores", "Último post",
           "Sinal encontrado", "Motivo", "E-mail", "WhatsApp", "Site",
           "Mensagem sugerida", "Perfil", "Busca de origem"]


def ler_perfis(caminho):
    perfis = []
    with open(caminho, encoding="utf-8") as f:
        for n, linha in enumerate(f, 1):
            linha = linha.strip()
            if not linha:
                continue
            try:
                perfis.append(json.loads(linha))
            except json.JSONDecodeError as e:
                print(f"Aviso: linha {n} ignorada (JSON inválido: {e})", file=sys.stderr)
    return perfis


def handle_de(p):
    h = (p.get("handle") or "").strip().lstrip("@").lower()
    if not h and p.get("url"):
        h = p["url"].rstrip("/").split("/")[-1].lower()
    return h


def nota(valor, maximo):
    try:
        return max(0, min(maximo, int(round(float(valor)))))
    except (TypeError, ValueError):
        return None


def atividade_pela_data(ultimo_post, hoje):
    try:
        d = datetime.strptime(str(ultimo_post)[:10], "%Y-%m-%d").date()
    except ValueError:
        return 0
    dias = (hoje - d).days
    if dias <= 15:
        return 20
    if dias <= 90:
        return 10
    return 0


def pontuar(p, hoje):
    ad = nota(p.get("aderencia"), 40) or 0
    si = nota(p.get("sinal"), 40) or 0
    at = nota(p.get("atividade"), 20)
    if at is None:
        at = atividade_pela_data(p.get("ultimo_post", ""), hoje)
    if not p.get("sinais"):
        si = 0  # regra da skill: sem evidência, sem sinal
    p["aderencia"], p["sinal"], p["atividade"] = ad, si, at
    p["pontuacao"] = ad + si + at
    return p


def mesclar(a, b):
    """Junta dois registros do mesmo @, preferindo campos preenchidos e somando evidências."""
    base, extra = (a, b) if a.get("pontuacao", 0) >= b.get("pontuacao", 0) else (b, a)
    for k, v in extra.items():
        if k in ("sinais", "pessoas"):
            vistos = {json.dumps(x, sort_keys=True) for x in base.get(k, [])}
            base.setdefault(k, [])
            base[k] += [x for x in v if json.dumps(x, sort_keys=True) not in vistos]
        elif not base.get(k) and v:
            base[k] = v
    return base


def handles_anteriores(caminho):
    if not caminho or not os.path.exists(caminho):
        return set()
    vistos = set()
    if caminho.lower().endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError:
            print("Aviso: openpyxl ausente; planilha anterior não foi lida.", file=sys.stderr)
            return set()
        wb = load_workbook(caminho, read_only=True)
        ws = wb["Leads"] if "Leads" in wb.sheetnames else wb.active
        linhas = ws.iter_rows(values_only=True)
        cab = [str(c or "").strip().lower() for c in next(linhas, [])]
        col = next((i for i, c in enumerate(cab) if c in ("@", "handle", "perfil @")), None)
        if col is None:
            return set()
        for r in linhas:
            if r and col < len(r) and r[col]:
                vistos.add(str(r[col]).strip().lstrip("@").lower())
    else:
        with open(caminho, encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                h = r.get("@") or r.get("handle") or ""
                if h:
                    vistos.add(h.strip().lstrip("@").lower())
    return vistos


def linha_lead(p):
    c = p.get("contato_publico") or {}
    sinal = p["sinais"][0]["texto"] if p.get("sinais") else ""
    url = p.get("url") or f"https://www.instagram.com/{p['handle']}/"
    return [p["pontuacao"], "@" + p["handle"], p.get("nome", ""), p.get("cidade", ""),
            p.get("seguidores", ""), p.get("ultimo_post", ""), sinal, p.get("motivo", ""),
            c.get("email", ""), c.get("whatsapp", ""), c.get("site", ""),
            p.get("mensagem", ""), url, p.get("fonte_busca", "")]


COMO_USAR = [
    "Prospecção via Instagram · trentinsolutions.com.br",
    "",
    "Pontuação (0 a 100) = aderência ao tema (até 40) + sinal de compra (até 40) + atividade recente (até 20).",
    "70 ou mais: abordar primeiro. 40 a 69: vale uma mensagem. Abaixo de 40: guardar para depois.",
    "A coluna Motivo traz a evidência encontrada no perfil, com data. Confira antes de abordar.",
    "Mensagem sugerida é só um rascunho. Ajuste ao seu jeito e envie você mesmo.",
    "Próxima rodada: passe esta planilha como --anterior para não repetir perfis.",
    "LGPD: dados públicos usados para contato comercial. Remova quem pedir para sair.",
    "",
    "Skill gratuita da Trentin Solutions, criada por Ivaldo Trentin · trentinsolutions.com.br",
    "Esta planilha te ajudou? Um Pix de qualquer valor em agradecimento é bem-vindo.",
    "Chave Pix aleatória: b73301d4-682f-46a0-857a-3a7653c2befb · Ivaldo Idalgo Trentin Junior Ltda",
]


def gravar_xlsx(leads, plano, saida):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    azul = PatternFill("solid", fgColor="3A6AB4")
    branco = Font(bold=True, color="FFFFFF", name="Arial")
    corpo = Font(name="Arial", size=10)
    faixas = [(70, "D8F3E6"), (40, "FCEFD6"), (0, "ECEFF4")]

    wb = Workbook()
    wb.properties.creator = "Trentin Solutions · Prospecção via Instagram"
    wb.properties.title = "Leads do Instagram"
    wb.properties.description = "Gerada pela skill gratuita Prospecção via Instagram, da Trentin Solutions (trentinsolutions.com.br)."
    wb.properties.keywords = "Trentin Solutions; prospecção; Instagram"
    ws = wb.active
    ws.title = "Leads"
    ws.append(COLUNAS)
    for p in leads:
        ws.append(linha_lead(p))
    larguras = [11, 24, 28, 16, 12, 13, 44, 52, 26, 18, 26, 52, 40, 26]
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cel in ws[1]:
        cel.fill, cel.font = azul, branco
        cel.alignment = Alignment(vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cel in row:
            cel.font = corpo
            cel.alignment = Alignment(vertical="top", wrap_text=cel.column in (7, 8, 12))
        nota_ = row[0].value or 0
        cor = next(c for lim, c in faixas if nota_ >= lim)
        row[0].fill = PatternFill("solid", fgColor=cor)
        row[0].font = Font(name="Arial", size=10, bold=True)
        link = row[12]
        if link.value:
            link.hyperlink = link.value
            link.font = Font(name="Arial", size=10, color="3A6AB4", underline="single")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    wp = wb.create_sheet("Pessoas")
    wp.append(["Perfil", "Pessoa", "Nome", "Papel", "Link"])
    for p in leads:
        for pe in p.get("pessoas", []):
            h = str(pe.get("handle", "")).lstrip("@")
            wp.append(["@" + p["handle"], "@" + h if h else "", pe.get("nome", ""),
                       pe.get("papel", ""), f"https://www.instagram.com/{h}/" if h else ""])
    for i, w in enumerate([24, 24, 28, 28, 40], 1):
        wp.column_dimensions[get_column_letter(i)].width = w
    for cel in wp[1]:
        cel.fill, cel.font = azul, branco
    wp.freeze_panes = "A2"

    wb_ = wb.create_sheet("Buscas")
    wb_.append(["Plano de busca usado"])
    wb_["A1"].fill, wb_["A1"].font = azul, branco
    for l in plano:
        wb_.append([l])
    wb_.column_dimensions["A"].width = 90

    wc = wb.create_sheet("Como usar")
    for l in COMO_USAR:
        wc.append([l])
    wc["A1"].font = Font(name="Arial", size=13, bold=True, color="3A6AB4")
    wc.column_dimensions["A"].width = 110
    wc.append([])
    wc.append([f"Gerado em {date.today().strftime('%d/%m/%Y')}"])
    wb.save(saida)


def gravar_csv(leads, saida):
    with open(saida, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(COLUNAS)
        for p in leads:
            w.writerow(linha_lead(p))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("entrada")
    ap.add_argument("--saida", default="leads_instagram.xlsx")
    ap.add_argument("--limite", type=int, default=0, help="quantos perfis manter (0 = todos)")
    ap.add_argument("--anterior", help="planilha de rodada anterior (.xlsx ou .csv)")
    ap.add_argument("--plano", default="trabalho/plano_buscas.md")
    a = ap.parse_args()

    hoje = date.today()
    brutos = ler_perfis(a.entrada)
    unicos = {}
    for p in brutos:
        h = handle_de(p)
        if not h:
            continue
        p["handle"] = h
        p = pontuar(p, hoje)
        unicos[h] = mesclar(unicos[h], p) if h in unicos else p

    ja_vistos = handles_anteriores(a.anterior)
    leads = [p for h, p in unicos.items() if h not in ja_vistos]
    leads.sort(key=lambda p: (p["pontuacao"], p.get("sinal", 0)), reverse=True)
    if a.limite:
        leads = leads[:a.limite]

    plano = []
    if a.plano and os.path.exists(a.plano):
        with open(a.plano, encoding="utf-8") as f:
            plano = [l.rstrip("\n") for l in f]

    saida = a.saida
    if saida.lower().endswith(".xlsx"):
        try:
            gravar_xlsx(leads, plano, saida)
        except ImportError:
            saida = saida[:-5] + ".csv"
            print("openpyxl não instalado; gravando CSV.", file=sys.stderr)
            gravar_csv(leads, saida)
    else:
        gravar_csv(leads, saida)

    fortes = sum(1 for p in leads if p.get("sinal", 0) >= 30)
    print(f"{len(brutos)} registros lidos · {len(unicos)} perfis únicos · "
          f"{len(unicos) - len([h for h in unicos if h not in ja_vistos])} já estavam na planilha anterior")
    print(f"{len(leads)} perfis na planilha · {fortes} com sinal forte · arquivo: {saida}")
    for p in leads[:5]:
        print(f"  {p['pontuacao']:>3}  @{p['handle']}  {p.get('motivo', '')}")


if __name__ == "__main__":
    main()
