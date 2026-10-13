#!/usr/bin/env python3
"""Seleciona produtos e gera o catálogo: página com pedido no WhatsApp, PDF, Status/Stories, carrossel e texto.

Trentin Solutions · criada por Ivaldo Trentin · trentinsolutions.com.br
Uso livre, mantendo este crédito (ver LICENSE.md).

  python catalogo.py selecionar produtos.xlsx --criterios margem,estoque --limite 12 [opções] --saida trabalho/selecao.json
  python catalogo.py gerar trabalho/selecao.json --marca trabalho/marca.json --fotos fotos/ \
         --formatos pagina,pdf,status,post,texto --saida catalogo/

Critérios (combináveis, na ordem de prioridade):
  margem      maior margem (preço − custo) ÷ preço
  estoque     mais unidades em estoque
  parados     estoque parado: tem estoque e não vende (ou não entra) há N dias; sugere preço promocional
  novidades   última compra/entrada mais recente
  vendidos    mais vendidos (precisa de coluna de vendas)
  todos       sem ranking, só os filtros
Filtros: --categorias "Velas,Canecas"  --preco-min 10 --preco-max 80  --codigos A1,A2 (escolha manual)
         --so-com-estoque  --margem-min 0.3 (margem mínima aceita, também protege a promoção)
"""
import argparse
import base64
import csv
import glob
import html
import io
import json
import math
import os
import re
import sys
import unicodedata
from datetime import date, datetime

# ---------------------------------------------------------------- leitura
ALIASES = {
    "codigo": ["codigo", "cod", "sku", "referencia", "ref"],
    "produto": ["produto", "nome", "descricao do produto", "item", "titulo"],
    "descricao": ["descricao", "detalhes", "observacao do catalogo", "texto"],
    "categoria": ["categoria", "tipo", "secao", "linha"],
    "preco": ["preco de venda", "preco venda", "preco", "valor de venda", "valor"],
    "preco_promo": ["preco promocional", "promocao", "preco promo", "por"],
    "custo": ["custo medio", "custo unit", "custo", "preco de custo"],
    "estoque": ["estoque atual", "estoque", "saldo", "qtd em estoque", "quantidade em estoque", "unidades em estoque"],
    "vendidos": ["qtd vendida", "vendidos", "vendas", "quantidade vendida"],
    "ultima_venda": ["ultima venda"],
    "ultima_compra": ["ultima compra", "ultima entrada", "data de entrada", "entrada"],
    "foto": ["foto", "imagem", "url da imagem", "url imagem", "link da foto", "arquivo da foto"],
    "status": ["status", "situacao", "ativo"],
}


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", s)).strip()


def slug(s):
    return norm(s).replace(" ", "-")


def mapear(cab):
    c = [norm(x) for x in cab]
    m = {}
    for campo, nomes in ALIASES.items():
        for nome in nomes:  # a ordem dos apelidos define a prioridade
            for i, x in enumerate(c):
                if i in m.values():
                    continue
                if x == nome or (len(nome) > 3 and x.startswith(nome)):
                    m.setdefault(campo, i)
            if campo in m:
                break
    return m


def numero(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("R$", "").replace("\u00a0", "").strip()
    if re.search(r",\d{1,2}$", s):
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(re.sub(r"[^\d.\-]", "", s))
    except ValueError:
        return None


def data_(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            pass
    return None


def ler_produtos(caminho, aba=None):
    if caminho.lower().endswith(".json"):
        with open(caminho, encoding="utf-8") as f:
            dados = json.load(f)
        linhas = [list(dados[0].keys())] + [list(d.values()) for d in dados] if dados else []
    elif caminho.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(caminho, data_only=True, read_only=True)
        nomes = {norm(s): s for s in wb.sheetnames}
        ws = wb[nomes.get(norm(aba or "Produtos"), wb.sheetnames[0])]
        linhas = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        with open(caminho, encoding="utf-8-sig") as f:
            amostra = f.read(4096)
            f.seek(0)
            sep = ";" if amostra.count(";") > amostra.count(",") else ","
            linhas = list(csv.reader(f, delimiter=sep))
    if not linhas:
        return [], {}
    m = mapear(linhas[0])
    prods = []
    for r in linhas[1:]:
        g = lambda k: r[m[k]] if k in m and m[k] < len(r) else None
        nome = str(g("produto") or "").strip()
        if not nome:
            continue
        p = {"codigo": str(g("codigo") or "").strip().lstrip("'"), "produto": nome,
             "descricao": str(g("descricao") or "").strip() if "descricao" in m and m["descricao"] != m.get("produto") else "",
             "categoria": str(g("categoria") or "").strip(), "preco": numero(g("preco")),
             "preco_promo": numero(g("preco_promo")), "custo": numero(g("custo")), "estoque": numero(g("estoque")),
             "vendidos": numero(g("vendidos")), "ultima_venda": data_(g("ultima_venda")),
             "ultima_compra": data_(g("ultima_compra")), "foto": str(g("foto") or "").strip(),
             "status": str(g("status") or "").strip()}
        prods.append(p)
    return prods, m


# ---------------------------------------------------------------- seleção
def preco_promocional(p, desconto, margem_min):
    """Preço com desconto, arredondado para baixo num final comum, sem ficar abaixo da margem mínima."""
    alvo = p["preco"] * (1 - desconto)
    piso = (p["custo"] / (1 - margem_min)) if p.get("custo") else 0
    alvo = max(alvo, piso)
    if alvo >= 10:
        cand = math.floor(alvo) - 0.10 if alvo - math.floor(alvo) < 0.9 else math.floor(alvo) + 0.90
        cand = round(cand, 2)
        if cand < piso:
            cand = round(math.ceil(piso * 10) / 10, 2)
    else:
        cand = round(math.floor(alvo * 10) / 10, 2)
        if cand < piso:
            cand = round(math.ceil(piso * 10) / 10, 2)
    return cand if cand < p["preco"] - 0.009 else None


def cmd_selecionar(a):
    prods, m = ler_produtos(a.planilha, a.aba)
    hoje = data_(a.hoje) or date.today()
    avisos = []
    tem = {k: k in m for k in ("custo", "estoque", "vendidos", "ultima_compra", "ultima_venda", "categoria")}
    criterios = [c.strip() for c in a.criterios.split(",") if c.strip()]
    falta = {"margem": "custo", "estoque": "estoque", "parados": "estoque", "novidades": "ultima_compra", "vendidos": "vendidos"}
    for c in list(criterios):
        if c in falta and not tem[falta[c]]:
            avisos.append(f"critério '{c}' ignorado: a planilha não tem coluna de {falta[c].replace('_', ' ')}")
            criterios.remove(c)
    base = []
    for p in prods:
        if p["preco"] is None or p["preco"] <= 0:
            continue
        if norm(p["status"]) in ("inativo", "fora de linha", "descontinuado", "nao", "oculto"):
            continue
        if a.codigos and p["codigo"] not in [c.strip() for c in a.codigos.split(",")]:
            continue
        if a.categorias and norm(p["categoria"]) not in [norm(c) for c in a.categorias.split(",")]:
            continue
        if a.preco_min is not None and p["preco"] < a.preco_min:
            continue
        if a.preco_max is not None and p["preco"] > a.preco_max:
            continue
        if (a.so_com_estoque or "parados" in criterios) and tem["estoque"] and not (p["estoque"] or 0) > 0:
            continue
        p["margem"] = round((p["preco"] - p["custo"]) / p["preco"], 4) if p.get("custo") else None
        if a.margem_min and p["margem"] is not None and p["margem"] < a.margem_min and "parados" not in criterios:
            continue
        ref = p["ultima_venda"] or p["ultima_compra"]
        p["dias_parado"] = (hoje - ref).days if ref else None
        base.append(p)

    def chave(p):
        k = []
        for c in criterios:
            if c == "margem":
                k.append(-(p["margem"] if p["margem"] is not None else -1))
            elif c == "estoque":
                k.append(-(p["estoque"] or 0))
            elif c == "parados":
                parado = (p["dias_parado"] or 0) >= a.dias_parado
                k.append((0 if parado else 1, -(p["estoque"] or 0) * (p["custo"] or p["preco"])))
            elif c == "novidades":
                k.append(-(p["ultima_compra"].toordinal() if p["ultima_compra"] else 0))
            elif c == "vendidos":
                k.append(-(p["vendidos"] or 0))
        return k

    if "parados" in criterios:
        base = [p for p in base if (p["dias_parado"] or 0) >= a.dias_parado or len(criterios) > 1]
    base.sort(key=chave)
    escolhidos = base[: a.limite] if a.limite else base

    for p in escolhidos:
        motivos = []
        if p["margem"] is not None:
            motivos.append(f"margem {p['margem']:.0%}")
        if p["estoque"] is not None:
            motivos.append(f"{p['estoque']:g} em estoque")
        if p["vendidos"] is not None and "vendidos" in criterios:
            motivos.append(f"{p['vendidos']:g} vendidos")
        if p["ultima_compra"] and "novidades" in criterios:
            motivos.append(f"chegou em {p['ultima_compra'].strftime('%d/%m')}")
        if "parados" in criterios and p["dias_parado"] is not None and p["dias_parado"] >= a.dias_parado:
            motivos.append(f"parado há {p['dias_parado']} dias")
            if not p.get("preco_promo"):
                promo = preco_promocional(p, a.desconto, a.margem_min or 0.2)
                if promo:
                    p["preco_promo"] = promo
                    p["promo_sugerida"] = True
                    motivos.append(f"promoção sugerida {brl(promo)} (de {brl(p['preco'])})")
                else:
                    motivos.append("sem espaço para desconto acima da margem mínima")
        p["motivo"] = " · ".join(motivos)
    saida = {"gerado_em": hoje.isoformat(), "fonte": os.path.basename(a.planilha), "criterios": criterios,
             "filtros": {"categorias": a.categorias, "preco_min": a.preco_min, "preco_max": a.preco_max,
                         "codigos": a.codigos, "margem_min": a.margem_min, "dias_parado": a.dias_parado},
             "avisos": avisos, "total_na_planilha": len(prods), "candidatos": len(base),
             "produtos": [{k: (v.isoformat() if isinstance(v, date) else v) for k, v in p.items()} for p in escolhidos]}
    os.makedirs(os.path.dirname(os.path.abspath(a.saida)), exist_ok=True)
    with open(a.saida, "w", encoding="utf-8") as f:
        json.dump(saida, f, ensure_ascii=False, indent=2)
    print(f"{len(prods)} produtos na planilha · {len(base)} passaram nos filtros · {len(escolhidos)} escolhidos -> {a.saida}")
    for av in avisos:
        print("  aviso:", av)
    for p in escolhidos:
        print(f"  {p['codigo'] or '-':<10} {p['produto'][:34]:<34} {brl(p['preco']):>11}  {p['motivo']}")


# ---------------------------------------------------------------- utilidades visuais
def brl(v):
    if v is None:
        return ""
    s = f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return "R$ " + s


def hex_rgb(h):
    h = h.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def luminancia(rgb):
    def c(x):
        x /= 255
        return x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4
    r, g, b = rgb
    return 0.2126 * c(r) + 0.7152 * c(g) + 0.0722 * c(b)


def texto_sobre(rgb):
    return (255, 255, 255) if luminancia(rgb) < 0.45 else (20, 22, 28)


def misturar(a, b, t):
    return tuple(int(a[i] + (b[i] - a[i]) * t) for i in range(3))


FONTES = {
    "titulo": ["Inter-Bold.otf", "InterDisplay-Bold.otf", "Inter-SemiBold.otf", "DejaVuSans-Bold.ttf",
               "LiberationSans-Bold.ttf", "Arial Bold.ttf", "arialbd.ttf", "Helvetica-Bold.ttf"],
    "texto": ["Inter-Regular.otf", "Inter-Medium.otf", "DejaVuSans.ttf", "LiberationSans-Regular.ttf",
              "Arial.ttf", "arial.ttf"],
}
_cache_fonte = {}


def achar_fonte(papel, preferida=None):
    if preferida and os.path.exists(preferida):
        return preferida
    if papel in _cache_fonte:
        return _cache_fonte[papel]
    pastas = ["/usr/share/fonts", "/usr/local/share/fonts", os.path.expanduser("~/.fonts"),
              "/Library/Fonts", "/System/Library/Fonts", "C:/Windows/Fonts"]
    for nome in FONTES[papel]:
        for p in pastas:
            achados = glob.glob(os.path.join(p, "**", nome), recursive=True)
            if achados:
                _cache_fonte[papel] = achados[0]
                return achados[0]
    _cache_fonte[papel] = None
    return None


def fonte(papel, tam, marca):
    from PIL import ImageFont
    caminho = achar_fonte(papel, marca.get("fonte_" + papel))
    if caminho:
        return ImageFont.truetype(caminho, tam)
    try:
        return ImageFont.load_default(size=tam)
    except TypeError:
        return ImageFont.load_default()


def quebrar(draw, texto, f, largura, max_linhas=3):
    palavras, linhas, atual = str(texto).split(), [], ""
    for w in palavras:
        teste = (atual + " " + w).strip()
        if draw.textlength(teste, font=f) <= largura:
            atual = teste
        else:
            if atual:
                linhas.append(atual)
            atual = w
    if atual:
        linhas.append(atual)
    if len(linhas) > max_linhas:
        linhas = linhas[:max_linhas]
        while linhas[-1] and draw.textlength(linhas[-1] + "…", font=f) > largura:
            linhas[-1] = linhas[-1][:-1]
        linhas[-1] = linhas[-1].rstrip() + "…"
    return linhas


def abrir_foto(p, pasta):
    from PIL import Image, ImageOps
    cands = []
    if p.get("foto"):
        cands.append(p["foto"] if os.path.isabs(p["foto"]) else os.path.join(pasta or ".", p["foto"]))
        cands.append(p["foto"])
    if pasta and os.path.isdir(pasta):
        arquivos = [f for f in os.listdir(pasta) if f.lower().endswith((".jpg", ".jpeg", ".png", ".webp"))]
        chaves = [x for x in (p.get("codigo"), slug(p["produto"])) if x]
        for f in arquivos:
            base = os.path.splitext(f)[0]
            if any(norm(base) == norm(k) or norm(base).startswith(norm(k) + " ") for k in chaves):
                cands.append(os.path.join(pasta, f))
    for c in cands:
        if c and os.path.isfile(c):
            try:
                return ImageOps.exif_transpose(Image.open(c)).convert("RGB")
            except Exception:  # noqa: BLE001
                pass
    return None


def cobrir(img, w, h):
    from PIL import ImageOps
    return ImageOps.fit(img, (w, h), centering=(0.5, 0.45))


def cartao_sem_foto(p, w, h, cor, marca):
    """Produto sem foto: cartão tipográfico com as iniciais, nunca imagem da internet."""
    from PIL import Image, ImageDraw
    fundo = misturar(cor, (255, 255, 255), 0.82)
    img = Image.new("RGB", (w, h), fundo)
    d = ImageDraw.Draw(img)
    ini = "".join(x[0] for x in norm(p["produto"]).split()[:2]).upper() or "?"
    f = fonte("titulo", int(h * 0.34), marca)
    tw = d.textlength(ini, font=f)
    d.text(((w - tw) / 2, h * 0.28), ini, font=f, fill=misturar(cor, fundo, 0.25))
    return img


def arredondar(img, raio):
    from PIL import Image, ImageDraw
    mask = Image.new("L", img.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle([0, 0, img.size[0] - 1, img.size[1] - 1], raio, fill=255)
    out = Image.new("RGBA", img.size)
    out.paste(img, (0, 0), mask)
    return out


def logo_img(marca, altura):
    from PIL import Image
    c = marca.get("logo")
    if c and os.path.isfile(c):
        im = Image.open(c).convert("RGBA")
        f = altura / im.height
        return im.resize((max(1, int(im.width * f)), altura), Image.LANCZOS)
    return None


# ---------------------------------------------------------------- formatos
def gerar_status(prods, marca, pasta_fotos, saida):
    from PIL import Image, ImageDraw
    cor = hex_rgb(marca.get("cor", "#3a6ab4"))
    cor2 = hex_rgb(marca.get("cor2", marca.get("cor", "#3a6ab4")))
    tc = texto_sobre(cor)
    arquivos = []
    for i, p in enumerate(prods, 1):
        W, H = 1080, 1920
        img = Image.new("RGB", (W, H), cor)
        d = ImageDraw.Draw(img)
        for y in range(H):  # degradê vertical suave
            d.line([(0, y), (W, y)], fill=misturar(cor, cor2, y / H * 0.6))
        lg = logo_img(marca, 90)
        if lg:
            img.paste(lg, (80, 110), lg)
        else:
            d.text((80, 120), marca.get("nome", ""), font=fonte("titulo", 52, marca), fill=tc)
        foto = abrir_foto(p, pasta_fotos)
        fw, fh = 920, 980
        fimg = cobrir(foto, fw, fh) if foto else cartao_sem_foto(p, fw, fh, cor, marca)
        img.paste(arredondar(fimg, 48), (80, 290), arredondar(fimg, 48))
        y = 1320
        if p.get("categoria"):
            d.text((80, y), p["categoria"].upper(), font=fonte("texto", 34, marca), fill=misturar(tc, cor, 0.35))
            y += 58
        for linha in quebrar(d, p["produto"], fonte("titulo", 68, marca), 920, 2):
            d.text((80, y), linha, font=fonte("titulo", 68, marca), fill=tc)
            y += 82
        y += 20
        if marca.get("mostrar_preco", True):
            if p.get("preco_promo"):
                fp = fonte("texto", 44, marca)
                de = "de " + brl(p["preco"])
                d.text((80, y + 18), de, font=fp, fill=misturar(tc, cor, 0.4))
                wde = d.textlength(de, font=fp)
                d.line([(80, y + 44), (80 + wde, y + 44)], fill=misturar(tc, cor, 0.4), width=4)
                d.text((80 + wde + 28, y), "por " + brl(p["preco_promo"]), font=fonte("titulo", 84, marca), fill=tc)
            else:
                d.text((80, y), brl(p["preco"]), font=fonte("titulo", 96, marca), fill=tc)
        pill = marca.get("chamada", "Peça pelo WhatsApp")
        fpill = fonte("titulo", 40, marca)
        pw = d.textlength(pill, font=fpill) + 96
        d.rounded_rectangle([80, 1740, 80 + pw, 1830], 45, fill=tc)
        d.text((80 + 48, 1762), pill, font=fpill, fill=cor)
        if marca.get("whatsapp_exibir"):
            d.text((80 + pw + 30, 1764), marca["whatsapp_exibir"], font=fonte("texto", 38, marca), fill=tc)
        nome = os.path.join(saida, f"status_{i:02d}_{slug(p['produto'])[:40]}.png")
        img.save(nome, optimize=True)
        arquivos.append(nome)
    return arquivos


def gerar_carrossel(prods, marca, pasta_fotos, saida, titulo):
    from PIL import Image, ImageDraw
    cor = hex_rgb(marca.get("cor", "#3a6ab4"))
    tc = texto_sobre(cor)
    fundo = (250, 249, 246)
    W, H = 1080, 1350
    arquivos = []
    # capa
    capa = Image.new("RGB", (W, H), cor)
    d = ImageDraw.Draw(capa)
    lg = logo_img(marca, 110)
    if lg:
        capa.paste(lg, (80, 90), lg)
    d.text((80, 260 if lg else 120), marca.get("nome", ""), font=fonte("texto", 46, marca), fill=tc)
    y = 360 if lg else 220
    for linha in quebrar(d, titulo, fonte("titulo", 110, marca), 920, 3):
        d.text((80, y), linha, font=fonte("titulo", 110, marca), fill=tc)
        y += 122
    mini = [p for p in prods][:3]
    x = 80
    for p in mini:
        foto = abrir_foto(p, pasta_fotos)
        t = cobrir(foto, 290, 380) if foto else cartao_sem_foto(p, 290, 380, cor, marca)
        capa.paste(arredondar(t, 28), (x, 820), arredondar(t, 28))
        x += 315
    d.text((80, 1240), "Arraste para ver →", font=fonte("texto", 38, marca), fill=tc)
    n = os.path.join(saida, "post_01_capa.png")
    capa.save(n, optimize=True)
    arquivos.append(n)
    # páginas 2×2
    for k in range(0, len(prods), 4):
        img = Image.new("RGB", (W, H), fundo)
        d = ImageDraw.Draw(img)
        for j, p in enumerate(prods[k:k + 4]):
            cx, cy = 60 + (j % 2) * 495, 60 + (j // 2) * 625
            foto = abrir_foto(p, pasta_fotos)
            t = cobrir(foto, 465, 405) if foto else cartao_sem_foto(p, 465, 405, cor, marca)
            img.paste(arredondar(t, 30), (cx, cy), arredondar(t, 30))
            yy = cy + 424
            for linha in quebrar(d, p["produto"], fonte("titulo", 36, marca), 465, 2):
                d.text((cx, yy), linha, font=fonte("titulo", 36, marca), fill=(20, 22, 28))
                yy += 44
            if marca.get("mostrar_preco", True):
                if p.get("preco_promo"):
                    fde = fonte("texto", 28, marca)
                    de = brl(p["preco"])
                    d.text((cx, yy + 12), de, font=fde, fill=(130, 130, 130))
                    wde = d.textlength(de, font=fde)
                    d.line([(cx, yy + 29), (cx + wde, yy + 29)], fill=(130, 130, 130), width=3)
                    d.text((cx + wde + 14, yy + 4), brl(p["preco_promo"]), font=fonte("titulo", 40, marca), fill=cor)
                else:
                    d.text((cx, yy + 4), brl(p["preco"]), font=fonte("titulo", 40, marca), fill=cor)
        d.text((60, H - 52), f"{marca.get('nome', '')} · {marca.get('chamada', 'Peça pelo WhatsApp')}",
               font=fonte("texto", 30, marca), fill=(110, 110, 110))
        n = os.path.join(saida, f"post_{k // 4 + 2:02d}.png")
        img.save(n, optimize=True)
        arquivos.append(n)
    return arquivos


def gerar_pdf(prods, marca, pasta_fotos, saida, titulo):
    from reportlab.lib.colors import Color
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    W, H = A4
    cor = hex_rgb(marca.get("cor", "#3a6ab4"))
    rc = Color(*[x / 255 for x in cor])
    tc = Color(*[x / 255 for x in texto_sobre(cor)])
    nome = os.path.join(saida, f"catalogo_{slug(marca.get('nome', 'loja'))}.pdf")
    c = canvas.Canvas(nome, pagesize=A4)
    c.setTitle(titulo)
    c.setAuthor(marca.get("nome", ""))
    c.setCreator("Trentin Solutions · Catálogo para WhatsApp e Instagram")
    por_pag, cols, mx, gap = 6, 2, 36, 16
    cw = (W - 2 * mx - gap) / cols
    ch, ih = 200, 136
    topo = H - 112
    wa = marca.get("whatsapp", "")
    paginas = max(1, math.ceil(len(prods) / por_pag))
    for pg in range(paginas):
        c.setFillColor(rc)
        c.rect(0, H - 92, W, 92, stroke=0, fill=1)
        c.setFillColor(tc)
        c.setFont("Helvetica-Bold", 22)
        c.drawString(mx, H - 52, marca.get("nome", ""))
        c.setFont("Helvetica", 11)
        c.drawString(mx, H - 72, titulo)
        for j, p in enumerate(prods[pg * por_pag:(pg + 1) * por_pag]):
            x = mx + (j % cols) * (cw + gap)
            yt = topo - (j // cols) * (ch + gap)
            foto = abrir_foto(p, pasta_fotos)
            px = (int(cw * 3), int(ih * 3))
            t = cobrir(foto, *px) if foto else cartao_sem_foto(p, px[0], px[1], cor, marca)
            buf = io.BytesIO()
            t.save(buf, "JPEG", quality=82)
            buf.seek(0)
            c.drawImage(ImageReader(buf), x, yt - ih, cw, ih)
            c.setFillColor(Color(0.1, 0.1, 0.12))
            c.setFont("Helvetica-Bold", 11.5)
            palavras, linha, linhas = p["produto"].split(), "", []
            for w_ in palavras:
                if c.stringWidth((linha + " " + w_).strip(), "Helvetica-Bold", 11.5) < cw:
                    linha = (linha + " " + w_).strip()
                else:
                    linhas.append(linha)
                    linha = w_
            linhas.append(linha)
            ty = yt - ih - 16
            for ln in linhas[:2]:
                c.drawString(x, ty, ln)
                ty -= 14
            py = yt - ch + 6
            if marca.get("mostrar_preco", True):
                if p.get("preco_promo"):
                    c.setFont("Helvetica", 10)
                    c.setFillColor(Color(0.5, 0.5, 0.5))
                    de = brl(p["preco"])
                    c.drawString(x, py, de)
                    wde = c.stringWidth(de, "Helvetica", 10)
                    c.line(x, py + 3.5, x + wde, py + 3.5)
                    c.setFillColor(rc)
                    c.setFont("Helvetica-Bold", 15)
                    c.drawString(x + wde + 8, py, brl(p["preco_promo"]))
                else:
                    c.setFillColor(rc)
                    c.setFont("Helvetica-Bold", 15)
                    c.drawString(x, py, brl(p["preco"]))
            if wa:
                msg = f"Oi! Vi no catálogo e quero: {p['produto']}"
                c.linkURL(f"https://wa.me/{wa}?text={quote(msg)}", (x, yt - ch, x + cw, yt), relative=0)
        c.setFillColor(Color(0.45, 0.45, 0.45))
        c.setFont("Helvetica", 9)
        rod = marca.get("chamada", "Peça pelo WhatsApp")
        if marca.get("whatsapp_exibir"):
            rod += " · " + marca["whatsapp_exibir"]
        if marca.get("validade"):
            rod += " · preços válidos até " + marca["validade"]
        c.drawString(mx, 24, rod)
        c.drawRightString(W - mx, 24, f"{pg + 1}/{paginas}")
        c.showPage()
    c.save()
    return [nome]


def quote(s):
    from urllib.parse import quote as q
    return q(s, safe="")


def gerar_texto(prods, marca, saida, titulo):
    linhas = [f"*{marca.get('nome', '')}* · {titulo}", ""]
    cats = {}
    for p in prods:
        cats.setdefault(p.get("categoria") or "Produtos", []).append(p)
    for cat, ps in cats.items():
        linhas.append(f"*{cat.upper()}*")
        for p in ps:
            if not marca.get("mostrar_preco", True):
                linhas.append(f"• {p['produto']}")
            elif p.get("preco_promo"):
                linhas.append(f"• {p['produto']} — ~{brl(p['preco'])}~ *{brl(p['preco_promo'])}*")
            else:
                linhas.append(f"• {p['produto']} — {brl(p['preco'])}")
        linhas.append("")
    if marca.get("validade"):
        linhas.append(f"_Preços válidos até {marca['validade']} ou enquanto durar o estoque._")
    linhas.append(marca.get("rodape_texto", "Para pedir, é só responder esta mensagem com o nome do produto e a quantidade."))
    nome = os.path.join(saida, "catalogo_whatsapp.txt")
    with open(nome, "w", encoding="utf-8") as f:
        f.write("\n".join(linhas) + "\n")
    return [nome]


def gerar_pagina(prods, marca, pasta_fotos, saida, titulo):
    cor = marca.get("cor", "#3a6ab4")
    rgb = hex_rgb(cor)
    sobre = "#ffffff" if texto_sobre(rgb) == (255, 255, 255) else "#14161c"
    itens = []
    for p in prods:
        foto = abrir_foto(p, pasta_fotos)
        t = cobrir(foto, 640, 640) if foto else cartao_sem_foto(p, 640, 640, rgb, marca)
        buf = io.BytesIO()
        t.save(buf, "JPEG", quality=78, optimize=True)
        itens.append({"id": p.get("codigo") or slug(p["produto"]), "nome": p["produto"], "cat": p.get("categoria") or "",
                      "desc": p.get("descricao") or "", "preco": p["preco"], "promo": p.get("preco_promo"),
                      "img": "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()})
    lg = ""
    if marca.get("logo") and os.path.isfile(marca["logo"]):
        with open(marca["logo"], "rb") as f:
            ext = os.path.splitext(marca["logo"])[1].lower().strip(".").replace("jpg", "jpeg")
            lg = f"data:image/{ext};base64," + base64.b64encode(f.read()).decode()
    dados = json.dumps({"itens": itens, "wa": marca.get("whatsapp", ""), "mostrarPreco": marca.get("mostrar_preco", True),
                        "loja": marca.get("nome", "")}, ensure_ascii=False)
    ld = {"@context": "https://schema.org", "@type": "ItemList", "name": f"{marca.get('nome', '')} · {titulo}",
          "itemListElement": [{"@type": "ListItem", "position": i + 1, "item": {
              "@type": "Product", "name": p["produto"], "category": p.get("categoria") or None,
              "offers": {"@type": "Offer", "priceCurrency": "BRL", "price": f"{(p.get('preco_promo') or p['preco']):.2f}",
                         "availability": "https://schema.org/InStock"}}} for i, p in enumerate(prods)]}
    pagina = PAGINA.replace("__TITULO__", html.escape(f"{marca.get('nome', '')} · {titulo}")) \
        .replace("__LOJA__", html.escape(marca.get("nome", ""))).replace("__SUB__", html.escape(titulo)) \
        .replace("__COR__", cor).replace("__SOBRE__", sobre) \
        .replace("__LOGO__", f'<img src="{lg}" alt="{html.escape(marca.get("nome", ""))}" class="logo">' if lg else "") \
        .replace("__VALIDADE__", html.escape(f"Preços válidos até {marca['validade']} ou enquanto durar o estoque." if marca.get("validade") else "")) \
        .replace("__DADOS__", dados.replace("</", "<\\/")).replace("__LD__", json.dumps(ld, ensure_ascii=False).replace("</", "<\\/")) \
        .replace("__INSTA__", html.escape(marca.get("instagram", "")))
    nome = os.path.join(saida, "catalogo.html")
    with open(nome, "w", encoding="utf-8") as f:
        f.write(pagina)
    return [nome]


PAGINA = r"""<!doctype html>
<html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>__TITULO__</title>
<script type="application/ld+json">__LD__</script>
<style>
:root{--cor:__COR__;--sobre:__SOBRE__;--fundo:#faf9f6;--tinta:#16181d;--suave:#6b6f78;--linha:rgba(22,24,29,.1)}
*{box-sizing:border-box}body{margin:0;background:var(--fundo);color:var(--tinta);font:16px/1.5 system-ui,-apple-system,"Segoe UI",Roboto,sans-serif}
header{background:var(--cor);color:var(--sobre);padding:28px 18px 22px}
.w{max-width:1040px;margin:0 auto}
.logo{height:48px;width:auto;display:block;margin-bottom:10px}
h1{margin:0;font-size:clamp(1.7rem,6vw,2.6rem);line-height:1.05;letter-spacing:-.02em}
header p{margin:6px 0 0;opacity:.85}
.busca{display:flex;gap:8px;flex-wrap:wrap;padding:14px 18px;position:sticky;top:0;background:var(--fundo);z-index:5;border-bottom:1px solid var(--linha)}
.busca input{flex:1;min-width:180px;border:1px solid var(--linha);border-radius:999px;padding:10px 16px;font:inherit;background:#fff}
.chips{display:flex;gap:6px;overflow-x:auto;width:100%;scrollbar-width:none}
.chip{border:1px solid var(--linha);background:#fff;border-radius:999px;padding:7px 13px;font:inherit;font-size:.9rem;white-space:nowrap;cursor:pointer}
.chip[aria-pressed=true]{background:var(--cor);color:var(--sobre);border-color:var(--cor)}
main{padding:18px 18px 120px}
.grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(160px,1fr));gap:14px}
.card{background:#fff;border:1px solid var(--linha);border-radius:18px;overflow:hidden;display:flex;flex-direction:column}
.card img{width:100%;aspect-ratio:1;object-fit:cover;display:block}
.card .c{padding:12px;display:flex;flex-direction:column;gap:4px;flex:1}
.card h2{font-size:.98rem;margin:0;line-height:1.25}
.card .d{font-size:.82rem;color:var(--suave);margin:0}
.preco{font-weight:700;font-size:1.08rem;color:var(--cor);margin-top:auto}
.preco s{color:var(--suave);font-weight:400;font-size:.85rem;margin-right:6px}
.qtd{display:flex;align-items:center;justify-content:space-between;margin-top:8px}
.qtd button{width:36px;height:36px;border-radius:50%;border:1px solid var(--linha);background:#fff;font-size:1.2rem;cursor:pointer}
.qtd button.add{background:var(--cor);color:var(--sobre);border-color:var(--cor)}
.qtd span{font-weight:700;min-width:2ch;text-align:center}
.barra{position:fixed;left:0;right:0;bottom:0;padding:12px 18px calc(12px + env(safe-area-inset-bottom));background:#fff;border-top:1px solid var(--linha);display:flex;align-items:center;gap:12px;justify-content:space-between;transform:translateY(110%);transition:transform .3s}
.barra.on{transform:none}
.barra a{background:#25D366;color:#fff;text-decoration:none;font-weight:700;padding:12px 18px;border-radius:999px;white-space:nowrap}
footer{color:var(--suave);font-size:.82rem;text-align:center;padding:24px 18px 110px}
footer a{color:inherit}
</style></head>
<body>
<header><div class="w">__LOGO__<h1>__LOJA__</h1><p>__SUB__</p></div></header>
<div class="busca"><div class="w" style="display:flex;flex-wrap:wrap;gap:10px;width:100%"><input id="q" type="search" placeholder="Buscar produto" aria-label="Buscar produto"><div class="chips" id="cats"></div></div></div>
<main><div class="w"><div class="grid" id="grid"></div></div></main>
<footer><div class="w"><p>__VALIDADE__</p><p>__INSTA__</p><p>Catálogo feito com a skill gratuita da <a href="https://trentinsolutions.com.br">Trentin Solutions</a></p></div></footer>
<div class="barra" id="barra"><div><b id="tot">R$ 0,00</b><div id="n" style="font-size:.85rem;color:var(--suave)"></div></div><a id="enviar" href="#" target="_blank" rel="noopener">Enviar pedido no WhatsApp</a></div>
<script>
const D=__DADOS__;const brl=v=>'R$ '+v.toFixed(2).replace('.',',').replace(/\B(?=(\d{3})+(?!\d))/g,'.');
const carr={};let cat='';const g=document.getElementById('grid');
const cats=[...new Set(D.itens.map(i=>i.cat).filter(Boolean))];
const cb=document.getElementById('cats');
if(cats.length>1){cb.innerHTML=['Tudo',...cats].map((c,i)=>`<button class="chip" aria-pressed="${i===0}" data-c="${i?c:''}">${c}</button>`).join('');
cb.onclick=e=>{const b=e.target.closest('.chip');if(!b)return;cat=b.dataset.c;[...cb.children].forEach(x=>x.setAttribute('aria-pressed',x===b));draw()}}
function draw(){const q=document.getElementById('q').value.toLowerCase();
g.innerHTML=D.itens.filter(i=>(!cat||i.cat===cat)&&(!q||(i.nome+' '+i.desc).toLowerCase().includes(q))).map(i=>`<article class="card"><img src="${i.img}" alt="${i.nome}" loading="lazy"><div class="c"><h2>${i.nome}</h2>${i.desc?`<p class="d">${i.desc}</p>`:''}${D.mostrarPreco?`<div class="preco">${i.promo?`<s>${brl(i.preco)}</s>${brl(i.promo)}`:brl(i.preco)}</div>`:''}<div class="qtd"><button aria-label="Tirar um" data-m="${i.id}">−</button><span>${carr[i.id]||0}</span><button class="add" aria-label="Pôr um" data-p="${i.id}">+</button></div></div></article>`).join('')}
g.onclick=e=>{const b=e.target.closest('button');if(!b)return;const id=b.dataset.p||b.dataset.m;carr[id]=Math.max(0,(carr[id]||0)+(b.dataset.p?1:-1));if(!carr[id])delete carr[id];draw();bar()};
document.getElementById('q').oninput=draw;
function bar(){const its=Object.entries(carr).map(([id,q])=>({i:D.itens.find(x=>x.id===id),q}));const n=its.reduce((s,x)=>s+x.q,0);const t=its.reduce((s,x)=>s+x.q*(x.i.promo||x.i.preco),0);
document.getElementById('barra').classList.toggle('on',n>0);document.getElementById('tot').textContent=D.mostrarPreco?brl(t):n+' itens';document.getElementById('n').textContent=n+(n===1?' item':' itens');
const msg='Oi! Quero fazer um pedido pelo catálogo:\n'+its.map(x=>`• ${x.q}× ${x.i.nome}${D.mostrarPreco?' ('+brl(x.i.promo||x.i.preco)+')':''}`).join('\n')+(D.mostrarPreco?`\nTotal: ${brl(t)}`:'');
document.getElementById('enviar').href=`https://wa.me/${D.wa}?text=${encodeURIComponent(msg)}`}
draw();
</script></body></html>"""


def cmd_gerar(a):
    with open(a.selecao, encoding="utf-8") as f:
        sel = json.load(f)
    with open(a.marca, encoding="utf-8") as f:
        marca = json.load(f)
    if marca.get("whatsapp"):
        marca["whatsapp"] = re.sub(r"\D", "", marca["whatsapp"])
        if len(marca["whatsapp"]) in (10, 11):
            marca["whatsapp"] = "55" + marca["whatsapp"]
    prods = sel["produtos"]
    if not prods:
        print("Nenhum produto na seleção.")
        return
    os.makedirs(a.saida, exist_ok=True)
    titulo = a.titulo or marca.get("titulo") or "Catálogo"
    formatos = [f.strip() for f in a.formatos.split(",")]
    feitos = []
    sem_foto = [p["produto"] for p in prods if abrir_foto(p, a.fotos) is None]
    if "pagina" in formatos:
        feitos += gerar_pagina(prods, marca, a.fotos, a.saida, titulo)
    if "pdf" in formatos:
        feitos += gerar_pdf(prods, marca, a.fotos, a.saida, titulo)
    if "status" in formatos:
        feitos += gerar_status(prods, marca, a.fotos, a.saida)
    if "post" in formatos:
        feitos += gerar_carrossel(prods, marca, a.fotos, a.saida, titulo)
    if "texto" in formatos:
        feitos += gerar_texto(prods, marca, a.saida, titulo)
    if not achar_fonte("titulo", marca.get("fonte_titulo")):
        print("Aviso: nenhuma fonte TrueType/OpenType encontrada; as imagens usaram a fonte padrão.")
    print(f"{len(prods)} produtos · {len(feitos)} arquivos em {a.saida}")
    for f_ in feitos:
        print("  ", os.path.basename(f_))
    if sem_foto:
        print(f"  {len(sem_foto)} sem foto (usei cartão com iniciais): " + ", ".join(sem_foto[:8]) + ("…" if len(sem_foto) > 8 else ""))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("selecionar")
    s.add_argument("planilha")
    s.add_argument("--aba")
    s.add_argument("--criterios", default="todos")
    s.add_argument("--limite", type=int, default=12)
    s.add_argument("--categorias")
    s.add_argument("--codigos")
    s.add_argument("--preco-min", type=float)
    s.add_argument("--preco-max", type=float)
    s.add_argument("--so-com-estoque", action="store_true")
    s.add_argument("--margem-min", type=float, default=0.0)
    s.add_argument("--dias-parado", type=int, default=60)
    s.add_argument("--desconto", type=float, default=0.15, help="desconto da promoção sugerida para parados")
    s.add_argument("--hoje", help="data de referência AAAA-MM-DD (padrão: hoje)")
    s.add_argument("--saida", default="trabalho/selecao.json")
    g = sub.add_parser("gerar")
    g.add_argument("selecao")
    g.add_argument("--marca", required=True)
    g.add_argument("--fotos")
    g.add_argument("--formatos", default="pagina,pdf,status,post,texto")
    g.add_argument("--titulo")
    g.add_argument("--saida", default="catalogo")
    a = ap.parse_args()
    {"selecionar": cmd_selecionar, "gerar": cmd_gerar}[a.cmd](a)


if __name__ == "__main__":
    main()
